import re
from datetime import date
from urllib.parse import parse_qs, urlparse

from fastapi.testclient import TestClient

from app.config import settings

from app.document_export import (
    DOCUMENT_LABELS,
    ExportData,
    build_document_body,
    build_docx_bytes,
    build_pdf_bytes,
    document_detail_table_data,
    guaranies_in_words,
)
from docx import Document as WordDocument
from pypdf import PdfReader
import io
from app.main import APP_BUILD, APP_VERSION, app


def login(client: TestClient) -> str:
    page = client.get("/login")
    match = re.search(r'name="csrf_token" value="([^"]+)"', page.text)
    assert page.status_code == 200 and match
    response = client.post(
        "/login",
        data={"email": "admin@digitlaboral.com.py", "password": settings.demo_admin_password, "csrf_token": match.group(1)},
        follow_redirects=False,
    )
    assert response.status_code == 303
    dashboard = client.get("/app")
    dashboard_token = re.search(r'name="csrf_token" value="([^"]+)"', dashboard.text)
    assert dashboard.status_code == 200 and dashboard_token
    return dashboard_token.group(1)


def test_app_version_and_build_are_visible_in_the_global_header():
    with TestClient(app) as client:
        login(client)
        response = client.get("/app")
        assert response.status_code == 200
        assert "app-version-indicator" in response.text
        assert f"v{APP_VERSION}" in response.text
        assert f"BUILD {APP_BUILD}" in response.text
        assert app.version == APP_VERSION



def test_certificate_a_preview_uses_model_text_and_uncached_static_asset():
    with TestClient(app) as client:
        login(client)
        page = client.get("/app/certificates")
        assert page.status_code == 200
        assert "no-store" in page.headers.get("cache-control", "")
        assert f"app.js?v={APP_VERSION}" in page.text
        script = client.get(f"/static/app.js?v={APP_VERSION}")
        assert script.status_code == 200
        assert "ha trabajado en esta firma comercial durante el lapso comprendido entre el " in script.text
        assert "se retira por su propia voluntad" in script.text
        assert "a nuestra entera satisfacción" in script.text
        assert "scheduleCanonicalCertificatePreview(type)" in script.text
        assert "formData.set('intent', 'preview')" in script.text
        assert "if (!companySelect.value || !employeeSelect.value) return;" in script.text
        assert "[NOMBRE DE LA EMPRESA]" in script.text
        assert "[NOMBRE COMPLETO DEL FUNCIONARIO]" in script.text
        assert "certificado_trabajo_b: 'Certificado de Trabajo B'" in script.text
        assert ".certificate-preview-signature-pair b,.certificate-preview-signature-pair small{display:block" in page.text


def test_certificate_b_matches_requested_model():
    title, body = build_document_body(
        "certificado_trabajo_b",
        company_name="KAMNIZ SRL",
        employee_name="DIEGO DAMIAN, FRANCO TABOADA",
        employee_document="6.638.398",
        position="DEPOSITERO",
        admission_date=date(2026, 7, 1),
        salary=0,
        issue_date=date(2026, 10, 29),
        city="Ciudad del Este",
        metadata={"effective_date": "2026-10-24", "last_period_salary_amount": 0},
    )
    assert title == "Certificado de Trabajo B"
    assert body.startswith("C E R T I F I C O:")
    assert "perteneció a la dotación de personal de esta firma KAMNIZ SRL" in body
    assert "habiendo ingresado a nuestro servicio el día 1 de julio de 2026" in body
    assert "prestado servicios hasta el día 24 de octubre de 2026" in body
    assert "desempeñándose como DEPOSITERO." in body
    assert "Este certificado lo expido a su solicitud y para su eventual uso ante terceros." in body
    assert "Salario devengado" not in body

    data = ExportData(
        document_type="certificado_trabajo_b",
        title=title,
        city="Ciudad del Este",
        issue_date=date(2026, 10, 29),
        company_name="KAMNIZ SRL",
        company_address="PIRIBEBUY E/MNHOR RODRIGUEZ",
        company_phone="(061)518-452",
        employee_name="DIEGO DAMIAN, FRANCO TABOADA",
        employee_document="6.638.398",
        position="DEPOSITERO",
        admission_date=date(2026, 7, 1),
        body=body,
        status="Emitido",
    )
    word = WordDocument(io.BytesIO(build_docx_bytes(data)))
    signature_text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    assert "KAMNIZ SRL\nPIRIBEBUY E/MNHOR RODRIGUEZ\nTEL. (061)518-452" in signature_text
    pdf = PdfReader(io.BytesIO(build_pdf_bytes(data)))
    pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "PIRIBEBUY E/MNHOR RODRIGUEZ" in pdf_text
    assert "TEL. (061)518-452" in pdf_text


def test_certificate_preview_endpoint_uses_same_body_as_export_model():
    with TestClient(app) as client:
        csrf = login(client)
        base_data = {
            "csrf_token": csrf,
            "company_id": 1,
            "employee_id": 1,
            "document_type": "certificado_trabajo_a",
            "city": "Ciudad del Este",
            "issue_date": "2026-10-29",
            "position": "Auxiliar",
            "admission_date": "2026-07-01",
            "salary": "3044000",
            "effective_date": "2026-10-24",
            "last_period_salary_amount": "",
            "intent": "preview",
        }
        response = client.post("/app/certificates", data=base_data, follow_redirects=False)
        assert response.status_code == 200
        assert "no-store" in response.headers.get("cache-control", "")
        payload = response.json()
        assert payload["document_type"] == "certificado_trabajo_a"
        assert payload["body"].startswith("C E R T I F I C O:")
        assert "ha trabajado en esta firma comercial" in payload["body"]

        base_data["document_type"] = "certificado_trabajo_b"
        base_data["position"] = "DEPOSITERO"
        b_response = client.post("/app/certificates", data=base_data, follow_redirects=False)
        assert b_response.status_code == 200
        b_payload = b_response.json()
        assert b_payload["document_type"] == "certificado_trabajo_b"
        assert b_payload["title"] == "Certificado de Trabajo B"
        assert "perteneció a la dotación de personal" in b_payload["body"]
        assert "habiendo ingresado a nuestro servicio" in b_payload["body"]
        assert "desempeñándose como DEPOSITERO." in b_payload["body"]


def test_amount_in_words_and_legal_basis():
    assert guaranies_in_words(2_739_341) == "dos millones setecientos treinta y nueve mil trescientos cuarenta y un guaraníes"
    title, body = build_document_body(
        "aguinaldo_anual",
        company_name="Empresa Demo S.A.",
        employee_name="Juan Pérez",
        employee_document="1.234.567",
        position="Auxiliar",
        admission_date=date(2024, 2, 1),
        salary=3_044_000,
        issue_date=date(2026, 12, 15),
        city="Ciudad del Este",
        metadata={
            "period_start": "2026-01-01",
            "period_end": "2026-12-31",
            "amount": 2_739_341,
            "notes": "",
        },
    )
    assert title == "Recibo de Pago de Aguinaldo"
    assert "artículo 243" in body
    assert "2.739.341" in body


def test_word_and_pdf_downloads():
    with TestClient(app) as client:
        csrf = login(client)
        generator_page = client.get("/app/certificates")
        assert generator_page.status_code == 200
        assert "certificateLastPeriodSalary" in generator_page.text
        assert "certificateSalaryBasic" in generator_page.text
        assert "Certificado de Trabajo" in generator_page.text
        assert "Referencias oficiales utilizadas" not in generator_page.text
        assert "notificacion_vacaciones" in generator_page.text
        certificate_menu = re.search(
            r'<aside class="certificate-type-menu"[^>]*>(.*?)</aside>',
            generator_page.text,
            re.DOTALL,
        )
        assert certificate_menu
        assert re.findall(
            r'data-certificate-type="([^"]+)">([^<]+)</button>',
            certificate_menu.group(1),
        ) == [
            ("certificado_trabajo_a", "Certificado de Trabajo A"),
            ("certificado_trabajo_b", "Certificado de Trabajo B"),
            ("constancia", "Constancia de Trabajo"),
            ("notificacion_preaviso", "Notificación de Preaviso"),
            ("renuncia", "Notificación de Renuncia"),
            ("despido", "Notificación de Despidos"),
            ("solicitud_vacacion", "Solicitud de Vacación"),
            ("notificacion_vacaciones", "Notificación de Vacación"),
            ("usufructo_vacaciones", "Usufructo de Vacaciones"),
            ("contrato_trabajo", "Contrato de Trabajo"),
        ]
        assert 'id="certificatePreview"' in generator_page.text
        certificate_form = re.search(
            r'<form\b(?=[^>]*id="certificateForm")[^>]*>(.*?)</form>',
            generator_page.text,
            re.DOTALL,
        )
        assert certificate_form
        # A single payment_date control serves aguinaldo receipts and payroll receipts.
        # Duplicate names made the later hidden input override the visible date.
        assert len(re.findall(r'name="payment_date"', certificate_form.group(1))) == 1
        assert 'data-specific="aguinaldo_anual aguinaldo_proporcional recibo_salario"' in certificate_form.group(1)
        response = client.post(
            "/app/certificates",
            data={
                "csrf_token": csrf,
                "company_id": 1,
                "employee_id": 1,
                "document_type": "constancia",
                "city": "Ciudad del Este",
                "issue_date": "2026-08-01",
                "position": "Vendedor",
                "admission_date": "2024-02-01",
                "salary": "3044000",
                "nationality": "paraguaya",
                "civil_status": "",
                "observations": "Presentación bancaria.",
                # Browsers submit empty strings for optional controls. These must
                # be normalized rather than rejected by date/integer validation.
                "effective_date": "",
                "calculation_id": "",
                "period_start": "",
                "period_end": "",
                "leave_start": "",
                "leave_end": "",
                "contract_end_date": "",
                "payment_date": "",
                "vacation_request_date": "",
                "vacation_notice_date": "",
                "vacation_payment_date": "",
                "birth_event_date": "",
                "last_worked_date": "",
                "notice_start_date": "",
                "incident_date": "",
                "sanction_start_date": "",
                "sanction_end_date": "",
                "intent": "save",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        query = parse_qs(urlparse(response.headers["location"]).query)
        certificate_id = int(query["created"][0])

        docx = client.get(f"/app/certificates/{certificate_id}/download.docx")
        assert docx.status_code == 200
        assert docx.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        assert docx.content[:2] == b"PK"

        pdf = client.get(f"/app/certificates/{certificate_id}/download.pdf")
        assert pdf.status_code == 200
        assert pdf.headers["content-type"].startswith("application/pdf")
        assert pdf.content.startswith(b"%PDF")



def test_work_certificate_uses_article_93_end_date_and_last_pay_period():
    title, body = build_document_body(
        "certificado_trabajo_a",
        company_name="Empresa Demo S.A.",
        employee_name="María González",
        employee_document="1.234.567",
        position="Auxiliar administrativa",
        admission_date=date(2022, 3, 1),
        salary=3_800_000,
        issue_date=date(2026, 10, 23),
        city="Ciudad del Este",
        metadata={
            "effective_date": "2026-10-15",
            "last_period_salary_amount": 3_550_000,
            "include_salary": False,
        },
    )
    assert title == "Certificado de Trabajo"
    assert "C E R T I F I C O:" in body
    assert "ha trabajado en esta firma comercial durante el lapso comprendido entre el 1 de marzo de 2022 y el 15 de octubre de 2026" in body
    assert "salario devengado durante el último período de pago" in body.lower()
    assert "Gs. 3.550.000" in body
    assert "se retira por su propia voluntad" in body
    assert "a nuestra entera satisfacción" in body
    assert "para su eventual uso ante terceros" in body
    assert "Remuneración mensual registrada: Gs. 3.800.000" not in body



def test_work_certificate_for_active_employee_does_not_claim_termination():
    title, body = build_document_body(
        "certificado_trabajo_a",
        company_name="Empresa Demo S.A.",
        employee_name="María González",
        employee_document="1.234.567",
        position="Auxiliar administrativa",
        admission_date=date(2022, 3, 1),
        salary=3_800_000,
        issue_date=date(2026, 10, 23),
        city="Ciudad del Este",
        metadata={"include_salary": False},
    )
    assert title == "Certificado de Trabajo"
    assert "C E R T I F I C O:" in body
    assert "presta servicios en la firma comercial Empresa Demo S.A. desde el 1 de marzo de 2022" in body
    assert "por su propia voluntad" not in body
    assert "eventual uso ante terceros" in body

def test_internal_contract_models_have_type_specific_official_model_fields():
    metadata = {
        "contract_modality": "mipyme_plazo_fijo",
        "contract_end_date": "2026-12-31",
        "contract_cause": "Necesidad temporal específica y documentada",
        "previous_contract_number": "CT-2026-01",
        "workplace": "Local comercial, Ciudad del Este",
        "job_functions": "Atención al cliente y caja",
        "work_schedule": "Lunes a viernes, 08:00 a 17:00",
        "weekly_rest": "Domingo",
        "payment_frequency": "Mensual",
        "probation_period": "Según procedencia legal",
    }
    types = (
        "contrato_trabajo",
        "contrato_tiempo_parcial",
        "contrato_mipymes_plazo_determinado",
        "contrato_mipymes_80",
        "adenda_prorroga_contrato",
    )
    for document_type in types:
        assert document_type in DOCUMENT_LABELS
        title, body = build_document_body(
            document_type,
            company_name="Empresa Demo S.A.",
            employee_name="María González",
            employee_document="1.234.567",
            position="Auxiliar administrativa",
            admission_date=date(2026, 1, 2),
            salary=3_800_000,
            issue_date=date(2026, 10, 23),
            city="Ciudad del Este",
            metadata=metadata,
            company_ruc="80000000-0",
            company_address="Centro, Ciudad del Este",
            legal_representative="José Pérez",
        )
        assert title == DOCUMENT_LABELS[document_type]
        assert "Atención al cliente y caja" in body
        assert "Lunes a viernes, 08:00 a 17:00" in body
        assert "BORRADOR" in body or "REVISAR" in body


def test_payroll_receipt_table_breaks_down_official_income_and_deductions():
    metadata = {
        "salary_basic_amount": 3_500_000,
        "overtime_day_amount": 150_000,
        "overtime_night_amount": 100_000,
        "ips_deduction_amount": 341_000,
        "advance_deduction_amount": 100_000,
    }
    detail = document_detail_table_data("recibo_salario", metadata, 3_500_000)
    assert detail is not None
    assert detail["headers"] == ["Concepto", "Haberes (Gs.)", "Descuentos (Gs.)"]
    text = " ".join(str(cell) for row in detail["rows"] for cell in row["cells"])
    assert "Horas extras diurnas (50%)" in text
    assert "Horas extras nocturnas (100%)" in text
    assert "TOTAL INGRESOS" in text
    assert "TOTAL DESCUENTOS" in text
    assert "NETO A PERCIBIR" in text
    assert "3.750.000" in text
    assert "441.000" in text
    assert "3.309.000" in text


def test_payroll_detail_table_appears_in_generated_word_and_pdf():
    data = ExportData(
        document_type="recibo_salario",
        title="Recibo de Salario",
        city="Ciudad del Este",
        issue_date=date(2026, 10, 23),
        company_name="Empresa Demo S.A.",
        company_ruc="80000000-0",
        employee_name="María González",
        employee_document="1.234.567",
        position="Auxiliar administrativa",
        admission_date=date(2024, 2, 1),
        salary=3_500_000,
        body="Liquidación de salario correspondiente al período declarado.",
        status="Emitido",
        metadata={
            "document_number": "DL-REC-0001",
            "period_start": "2026-10-01",
            "period_end": "2026-10-31",
            "payment_date": "2026-10-31",
            "payment_method": "Transferencia bancaria",
            "days_worked": "30",
            "hours_worked": "240",
            "salary_basic_amount": 3_500_000,
            "overtime_day_amount": 150_000,
            "ips_deduction_amount": 341_000,
            "advance_deduction_amount": 100_000,
        },
    )
    word = WordDocument(io.BytesIO(build_docx_bytes(data)))
    word_text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    word_text += "\n" + "\n".join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
    assert "TOTAL INGRESOS" in word_text
    assert "NETO A PERCIBIR" in word_text

    pdf = PdfReader(io.BytesIO(build_pdf_bytes(data)))
    pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "TOTAL INGRESOS" in pdf_text
    assert "NETO A PERCIBIR" in pdf_text



def test_work_constancia_uses_last_pay_period_for_termination():
    title, body = build_document_body(
        "constancia",
        company_name="Empresa Demo S.A.",
        employee_name="María González",
        employee_document="1.234.567",
        position="Auxiliar administrativa",
        admission_date=date(2022, 3, 1),
        salary=3_800_000,
        issue_date=date(2026, 10, 23),
        city="Ciudad del Este",
        metadata={
            "effective_date": "2026-10-15",
            "last_period_salary_amount": 3_550_000,
            "include_salary": "off",
        },
    )
    assert title == "Constancia de Trabajo"
    assert "15 de octubre de 2026" in body
    assert "Salario devengado durante el último período de pago" in body
    assert "3.550.000" in body
    assert "3.800.000" not in body
