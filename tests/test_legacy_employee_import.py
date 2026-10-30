from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook

from app.legacy_employee_import import parse_legacy_employee_workbook


def _book_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BASE DE DATOS TRABAJADORES"
    headers = [
        "CODTRA", "NOMTRA", "APETRA", "CODEMP", "CODSEC", "CEDIDE",
        "ENTMAN", "SALMAN", "ENTTAR", "SALTAR", "ESTCIV", "NACION",
        "FECNAC", "LUGNAC", "PROFES", "DOMACT", "SEXOXX", "CIUDAD",
        "TELEFO", "NUMIPS", "IPSXXX", "OCUPAC", "CANHIJ", "SALACT",
        "FECENT", "FECSAL",
    ]
    sheet.append(headers)
    sheet.append([
        "10001", "ANA", "PEREZ", "1001", "001", "1234567",
        "08:00", "12:00", "13:00", "17:00", "S", "Paraguaya",
        datetime(2000, 1, 1), "CDE", "Comerciante", "Centro", "F", "CDE",
        "061000000", "12345", "S", "Vendedora", "1", 1500000,
        datetime(2020, 1, 2), datetime(2021, 1, 2),
    ])
    sheet.append([
        "10002", "ANA", "PEREZ", "1001", "001", "1234567",
        "08:00", "12:00", "13:00", "17:00", "S", "Paraguaya",
        datetime(2025, 1, 1), "CDE", "Comerciante", "Centro", "F", "CDE",
        "061000000", "12345", "S", "Vendedora", "1", 2500000,
        datetime(2022, 2, 3), None,
    ])
    sheet.append([
        "10003", "LUIS", "VERA", "1001", "001", "0",
        None, None, None, None, "C", "Paraguaya",
        None, None, "Operario", "Barrio", "M", "CDE",
        None, None, "N", "Operario", "0", 1800000,
        datetime(2019, 3, 1), datetime(2019, 12, 1),
    ])
    sheet.append([
        "10004", None, None, "1001", "001", "1234568",
        None, None, None, None, None, "Paraguaya",
        None, None, None, None, None, None, None, None, "N",
        "Auxiliar", "0", 1500000, datetime(2018, 1, 1), datetime(2018, 6, 1),
    ])
    sheet.append([
        "10005", "MARTA", "GOMEZ", "1001", "001", "1234569",
        None, None, None, None, None, "Paraguaya",
        None, None, "Comerciante", None, "F", "CDE",
        None, None, "N", "Vendedora", "0", 1500000, None, datetime(2018, 6, 1),
    ])
    sheet.append([
        "10006", "CARLOS", "DIAZ", "1002", "001", "9988776",
        None, None, None, None, None, "Paraguaya",
        None, None, "Operario", None, "M", "CDE",
        None, None, "N", "Operario", "0", 2000000,
        datetime(2017, 5, 1), datetime(2018, 5, 1),
    ])
    sheet.append([
        "10007", "JOSE", "LOPEZ", "1001", "001", "7777777",
        None, None, None, None, None, "Paraguaya",
        None, None, "Operario", None, "M", "CDE",
        None, None, "N", "Operario", "0", 2000000,
        datetime(2020, 5, 1), datetime(2019, 5, 1),
    ])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_legacy_import_groups_rehires_and_preserves_history_without_duplicate_people():
    """Consolidate rehires and reject untrustworthy dates without losing the timeline."""
    company_map = {
        "1001": [{
            "id": 11, "name": "Empresa Uno", "status": "Activa",
            "branch_id": 21, "department": "Alto Paraná",
        }],
        "1002": [
            {"id": 12, "name": "Empresa Dos", "status": "Activa", "branch_id": 22, "department": ""},
            {"id": 13, "name": "Otra Empresa Dos", "status": "Activa", "branch_id": 23, "department": ""},
        ],
    }
    result = parse_legacy_employee_workbook(
        _book_bytes(), "trabajadores.xlsx", company_map, today=date(2026, 10, 30)
    )

    assert result["summary"]["source_rows"] == 7
    assert result["summary"]["identity_groups"] == 6
    assert result["summary"]["importable"] == 3
    assert result["summary"]["excluded"] == 3
    assert result["summary"]["active"] == 1
    assert result["summary"]["terminated"] == 2
    assert result["summary"]["invalid_termination_dates"] == 1
    assert result["summary"]["repeated_identity_groups"] == 1
    assert result["summary"]["without_cedula"] == 1

    ana = next(item for item in result["candidates"] if item["full_name"] == "ANA PEREZ")
    assert ana["document_number"] == "1234567"
    assert ana["status"] == "Activo"
    assert ana["admission_date"] == "2022-02-03"
    assert ana["base_salary"] == 2500000
    assert ana["termination_date"] is None
    assert [event["event_type"] for event in ana["events"]] == [
        "Ingreso histórico", "Egreso histórico", "Reingreso histórico"
    ]
    assert ana["profile"]["work_schedule"].startswith("Mañana: 08:00–12:00")

    luis = next(item for item in result["candidates"] if item["full_name"] == "LUIS VERA")
    assert luis["document_number"] == "S/D-1001-10003"
    assert luis["synthetic_document"] is True
    assert luis["status"] == "Desvinculado"
    assert luis["termination_date"] == "2019-12-01"

    jose = next(item for item in result["candidates"] if item["full_name"] == "JOSE LOPEZ")
    assert jose["status"] == "Desvinculado"
    assert jose["termination_date"] is None
    assert "fechas de egreso" in jose["notes"]

    reasons = {item["reason"] for item in result["excluded"]}
    assert any("nombre" in reason.lower() for reason in reasons)
    assert any("fecha de ingreso" in reason.lower() for reason in reasons)
    assert any("código patronal" in reason.lower() for reason in reasons)


def test_legacy_import_rejects_non_xlsx_and_missing_company_mapping():
    try:
        parse_legacy_employee_workbook(b"not an xlsx", "trabajadores.csv", {})
    except ValueError as exc:
        assert ".xlsx" in str(exc)
    else:
        raise AssertionError("CSV must not be accepted by the legacy workbook importer")


def _single_company_map():
    return {"1001": [{"id": 11, "name": "Empresa Uno", "status": "Activa"}]}


def test_historical_policy_excludes_bad_exits_and_never_activates_open_source_records():
    from app.legacy_employee_import import validate_legacy_employee_policy

    result = parse_legacy_employee_workbook(
        _book_bytes(), "trabajadores.xlsx", _single_company_map(),
        historical_only=True, exclude_inconsistent_terminations=True,
    )
    assert result["summary"]["importable"] == 2
    assert result["summary"]["excluded"] == 4
    assert result["summary"]["excluded_inconsistent_terminations"] == 1
    assert result["summary"]["invalid_termination_dates"] == 0
    assert result["summary"]["will_activate"] == 0
    assert result["summary"]["will_terminate"] == 2
    assert result["summary"]["historical_review"] == 1
    ana = next(item for item in result["candidates"] if item["full_name"] == "ANA PEREZ")
    assert ana["active_in_source"] is True
    assert ana["status"] == "Desvinculado"
    assert ana["termination_date"] is None
    assert "pendiente de revisión" in ana["notes"]
    assert [event["event_type"] for event in ana["events"]] == [
        "Ingreso histórico", "Egreso histórico", "Reingreso histórico"
    ]
    validate_legacy_employee_policy(result)


def test_excluding_bad_exits_removes_the_whole_identity_even_with_a_later_open_rehire():
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(_book_bytes()))
    workbook.active.cell(2, 26, datetime(2019, 1, 1))
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    result = parse_legacy_employee_workbook(
        content.getvalue(), "trabajadores.xlsx", _single_company_map(),
        historical_only=True, exclude_inconsistent_terminations=True,
    )
    assert [item["full_name"] for item in result["candidates"]] == ["LUIS VERA"]
    ana = next(item for item in result["excluded"] if item["name"] == "ANA PEREZ")
    assert ana["source_rows"] == [2, 3]
    assert ana["exclusion_type"] == "inconsistent_termination"


def test_confirmation_policy_rejects_active_or_inconsistent_candidates_before_writing():
    from copy import deepcopy
    import pytest
    from app.legacy_employee_import import validate_legacy_employee_policy

    result = parse_legacy_employee_workbook(
        _book_bytes(), "trabajadores.xlsx", _single_company_map(),
        historical_only=True, exclude_inconsistent_terminations=True,
    )
    active = deepcopy(result)
    active["candidates"][0]["status"] = "Activo"
    with pytest.raises(ValueError, match="legajo activo"):
        validate_legacy_employee_policy(active)
    inconsistent = deepcopy(result)
    inconsistent["candidates"][0]["warning_invalid_termination"] = True
    with pytest.raises(ValueError, match="egresos excluidos"):
        validate_legacy_employee_policy(inconsistent)


def test_historical_upload_and_confirmation_preserve_existing_active_employee_and_create_no_settlement(tmp_path, monkeypatch):
    import re
    import uuid
    from fastapi.testclient import TestClient
    from sqlalchemy import select, func
    import app.main as main
    from app.config import settings
    from app.database import SessionLocal
    from app.models import Company, Employee, EmployeeComplianceProfile, EmployeeEvent, EmployeeExit, FinalSettlement, User

    monkeypatch.setattr(main, "IMPORT_DIR", tmp_path)
    suffix = uuid.uuid4().hex[:12]
    company_id = None
    with TestClient(main.app) as client:
        login_page = client.get("/login")
        csrf = re.search(r'name="csrf_token" value="([^"]+)"', login_page.text).group(1)
        login = client.post("/login", data={
            "email": "admin@digitlaboral.com.py", "password": settings.demo_admin_password,
            "csrf_token": csrf,
        }, follow_redirects=False)
        assert login.status_code == 303
        dashboard = client.get("/app")
        csrf = re.search(r'name="csrf_token" value="([^"]+)"', dashboard.text).group(1)
        with SessionLocal() as db:
            studio_id = db.scalar(select(User.studio_id).where(User.email == "admin@digitlaboral.com.py"))
            company = Company(
                studio_id=studio_id, company_code=int(uuid.uuid4().int % 1_000_000_000) + 1,
                legacy_code="QA-" + suffix, legal_name="Empresa histórica QA " + suffix,
                ruc="QA-" + suffix, status="Activa",
            )
            db.add(company)
            db.flush()
            company_id = company.id
            db.add(Employee(
                company_id=company_id, employee_code=20, full_name="Existente activo",
                document_number="EX-" + suffix, position="Auxiliar", status="Activo",
                admission_date=date(2024, 1, 1), base_salary=4_000_000,
            ))
            db.commit()
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["CODTRA", "NOMTRA", "APETRA", "CODEMP", "CEDIDE", "FECENT", "FECSAL", "SALACT"])
        sheet.append([1, "Nueva", "Histórica", "QA-" + suffix, "NEW-" + suffix, date(2020, 1, 1), None, 3_000_000])
        sheet.append([2, "Existente", "Fuente", "QA-" + suffix, "EX-" + suffix, date(2018, 1, 1), date(2019, 1, 1), 1_000_000])
        sheet.append([3, "Egreso", "Inválido", "QA-" + suffix, "BAD-" + suffix, date(2020, 1, 1), date(2019, 1, 1), 1_000_000])
        sheet.append([4, None, None, "QA-" + suffix, "NONAME-" + suffix, date(2020, 1, 1), None, 1_000_000])
        sheet.append([5, "Sin", "Ingreso", "QA-" + suffix, "NODATE-" + suffix, None, None, 1_000_000])
        content = BytesIO()
        workbook.save(content)
        workbook.close()
        try:
            for attempt in range(2):
                preview = client.post(
                    "/app/import/legacy-employees/preview",
                    files={"file": ("historicos.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                    data={"csrf_token": csrf, "historical_only": "true", "exclude_inconsistent_terminations": "true"},
                    follow_redirects=False,
                )
                assert preview.status_code == 303, preview.text
                token = preview.headers["location"].rsplit("/", 1)[-1]
                stored = main._load_legacy_employee_preview(token)
                assert stored["summary"]["importable"] == 2
                assert stored["summary"]["excluded"] == 3
                assert stored["summary"]["will_activate"] == 0
                assert stored["options"] == {"historical_only": True, "exclude_inconsistent_terminations": True}
                shown = client.get(preview.headers["location"])
                assert shown.status_code == 200
                assert "Archivo histórico: todos como Desvinculado" in shown.text
                confirmed = client.post(preview.headers["location"] + "/confirm", data={"csrf_token": csrf}, follow_redirects=False)
                assert confirmed.status_code == 303, confirmed.text
                assert "legacy_imported=" + str(1 if attempt == 0 else 0) in confirmed.headers["location"]
                directory = client.get("/api/employees", params={"company_id": company_id, "include_archived": "true"})
                assert directory.status_code == 200 and directory.json()["count"] == 2
                with SessionLocal() as db:
                    employees = list(db.scalars(select(Employee).where(Employee.company_id == company_id)))
                    assert len(employees) == 2
                    active = next(item for item in employees if item.document_number == "EX-" + suffix)
                    assert active.status == "Activo" and active.base_salary == 4_000_000
                    assert active.admission_date == date(2024, 1, 1)
                    historical = next(item for item in employees if item.document_number == "NEW-" + suffix)
                    assert historical.status == "Desvinculado" and historical.termination_date is None
                    assert "pendiente de revisión" in historical.notes
                    assert db.scalar(select(func.count(EmployeeEvent.id)).where(EmployeeEvent.employee_id == historical.id)) == 1
                    assert db.scalar(select(func.count(EmployeeExit.id)).where(EmployeeExit.employee_id == historical.id)) == 0
                    assert db.scalar(select(func.count(FinalSettlement.id)).where(FinalSettlement.employee_id == historical.id)) == 0
            with SessionLocal() as db:
                db.add_all([
                    Employee(
                        company_id=company_id, employee_code=1000 + index,
                        full_name=f"ZZ Histórica QA {index:03d}", document_number=f"BULK-{suffix}-{index}",
                        position="Auxiliar", status="Desvinculado", admission_date=date(2020, 1, 1), base_salary=1_000_000,
                    ) for index in range(102)
                ])
                db.commit()
            filters = {"company_id": company_id, "status_filter": "Desvinculado", "show_archived": "true", "q": "ZZ Histórica QA"}
            first = client.get("/app/employees", params={**filters, "page": 1})
            second = client.get("/app/employees", params={**filters, "page": 2})
            assert first.status_code == second.status_code == 200
            assert first.text.count("<tr data-employee-row") == 100
            assert second.text.count("<tr data-employee-row") == 2
            assert "Página 2 de 2" in second.text
            assert "page=2" in first.text and "status_filter=Desvinculado" in first.text
            assert "ZZ Histórica QA 000" in first.text and "ZZ Histórica QA 000" not in second.text
            assert "ZZ Histórica QA 101" in second.text
        finally:
            with SessionLocal() as db:
                for employee in db.scalars(select(Employee).where(Employee.company_id == company_id)).all():
                    for event in db.scalars(select(EmployeeEvent).where(EmployeeEvent.employee_id == employee.id)).all():
                        db.delete(event)
                    for profile in db.scalars(select(EmployeeComplianceProfile).where(EmployeeComplianceProfile.employee_id == employee.id)).all():
                        db.delete(profile)
                    db.delete(employee)
                company = db.get(Company, company_id)
                if company:
                    db.delete(company)
                db.commit()
