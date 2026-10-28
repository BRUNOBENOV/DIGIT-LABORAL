from __future__ import annotations

import re
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.config import settings
from app.main import app
from app.database import SessionLocal
from app.models import CalculationRecord, Company
from app.report_center import REPORTS, get_report, make_xlsx, report_columns


def _login(client: TestClient) -> None:
    page = client.get("/login")
    token_match = re.search(r'name="csrf_token" value="([^"]+)"', page.text)
    assert page.status_code == 200 and token_match
    response = client.post(
        "/login",
        data={
            "email": "admin@digitlaboral.com.py",
            "password": settings.demo_admin_password,
            "csrf_token": token_match.group(1),
        },
        follow_redirects=False,
    )
    assert response.status_code == 303


def test_reports_catalog_lists_every_requested_submodule():
    assert len(REPORTS) == 35
    assert len(set(REPORTS)) == len(REPORTS)
    with TestClient(app) as client:
        _login(client)
        catalog = client.get("/api/reports/catalog")
        assert catalog.status_code == 200
        payload = catalog.json()
        keys = [report["key"] for group in payload["groups"] for report in group["reports"]]
        assert payload["total"] == 35
        assert len(keys) == 35
        assert set(keys) == set(REPORTS)
        page = client.get("/app/reports")
        assert page.status_code == 200
        assert "Informes laborales" in page.text
        assert "Listado Examen Médico" in page.text
        assert "Liquidación Salario" in page.text


def test_all_report_data_endpoints_return_paginated_schema():
    with TestClient(app) as client:
        _login(client)
        for key, spec in REPORTS.items():
            response = client.get(f"/api/reports/{key}?page=1&page_size=10")
            assert response.status_code == 200, f"{key}: {response.text[:300]}"
            payload = response.json()
            assert payload["report"]["key"] == key
            assert payload["report"]["label"] == spec.label
            assert isinstance(payload["records"], list)
            assert isinstance(payload["columns"], list) and payload["columns"]
            assert payload["page"] == 1
            assert payload["page_size"] == 10
            assert payload["total"] >= 0
            assert payload["pages"] >= 0
        invalid = client.get("/api/reports/no-existe")
        assert invalid.status_code == 404


def test_report_detail_and_excel_pdf_exports_work():
    with TestClient(app) as client:
        _login(client)
        detail = client.get("/app/reports/listado-empresa")
        assert detail.status_code == 200
        assert "Listado Empresa" in detail.text
        assert "Aplicar" in detail.text

        excel = client.get("/api/reports/listado-empresa/export?format=xlsx")
        assert excel.status_code == 200
        assert excel.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert excel.content.startswith(b"PK")

        pdf = client.get("/app/reports/listado-empresa/export?format=pdf")
        assert pdf.status_code == 200
        assert pdf.headers["content-type"].startswith("application/pdf")
        assert pdf.content.startswith(b"%PDF")

        bad_format = client.get("/api/reports/listado-empresa/export?format=exe")
        assert bad_format.status_code == 400


def test_report_queries_include_professional_suite_calculations():
    marker = "report-center-suite-data-check"
    with TestClient(app) as client:
        _login(client)
        with SessionLocal() as db:
            company = db.scalar(select(Company).order_by(Company.id))
            assert company is not None
            company_id = company.id
            for calc_type in ("suite_salary", "suite_settlement", "suite_aguinaldo", "suite_vacation"):
                db.add(
                    CalculationRecord(
                        company_id=company_id,
                        calculation_type=calc_type,
                        reference_period="2026-10",
                        amount=987654,
                        status="Borrador",
                        source="test",
                        notes=marker,
                        created_by="qa-report-center",
                    )
                )
            db.commit()

        cases = (
            ("liquidacion-salario", "suite_salary"),
            ("liquidacion-haberes", "suite_settlement"),
            ("liquidacion-aguinaldo", "suite_aguinaldo"),
            ("liquidacion-vacacion", "suite_vacation"),
        )
        expected_labels = {
            "suite_salary": "Salario neto",
            "suite_settlement": "Liquidación de haberes",
            "suite_aguinaldo": "Aguinaldo",
            "suite_vacation": "Vacaciones",
        }
        for report_key, expected_type in cases:
            response = client.get(
                f"/api/reports/{report_key}",
                params={"company_id": company_id, "q": marker, "page_size": 100},
            )
            assert response.status_code == 200, f"{report_key}: {response.text[:250]}"
            payload = response.json()
            assert any(
                row.get("Importe (Gs.)") == 987654
                and row.get("Tipo de cálculo") == expected_labels[expected_type]
                for row in payload["records"]
            ), report_key


def test_excel_export_forces_user_supplied_strings_to_text():
    spec = get_report("listado-empresa")
    assert spec is not None
    row = {column: "" for column in report_columns(spec)}
    row["Empresa"] = '=HYPERLINK("https://example.invalid","click")'
    row["Nombre comercial"] = '\t=HYPERLINK("https://example.invalid","click")'
    payload = make_xlsx(spec, [row])
    workbook = load_workbook(BytesIO(payload), data_only=False)
    cell = workbook["Informe"].cell(row=3, column=2)
    tab_prefixed = workbook["Informe"].cell(row=3, column=3)
    assert cell.data_type != "f"
    assert isinstance(cell.value, str) and cell.value.startswith("'=")
    assert tab_prefixed.data_type != "f"
    assert isinstance(tab_prefixed.value, str) and tab_prefixed.value.startswith("'\t=")
