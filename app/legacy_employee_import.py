from __future__ import annotations

import io
import math
import re
from collections import defaultdict
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


REQUIRED_COLUMNS = {"CODTRA", "CODEMP", "CEDIDE", "NOMTRA", "APETRA", "FECENT", "FECSAL"}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        if float(value).is_integer():
            return str(int(value))
    return str(value).strip()


def _date(value: Any, epoch: datetime | None) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        try:
            parsed = from_excel(value, epoch=epoch)
            return parsed.date() if isinstance(parsed, datetime) else parsed
        except (TypeError, ValueError, OverflowError):
            return None
    raw = _text(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None


def _digits(value: Any) -> str:
    raw = _text(value)
    digits = re.sub(r"\D", "", raw)
    if digits and set(digits) == {"0"}:
        return ""
    return digits


def _money(value: Any) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if not math.isfinite(float(value)) or float(value) < 0 or not float(value).is_integer():
            return None
        return int(value)
    raw = _text(value).replace("Gs.", "").replace("Gs", "").replace("₲", "").strip()
    raw = re.sub(r"\s+", "", raw)
    if not raw or raw.startswith("-"):
        return None
    if not re.fullmatch(r"(?:\d+|\d{1,3}(?:[.,]\d{3})+)", raw):
        return None
    return int(re.sub(r"[.,]", "", raw))


def _normal_ci(value: Any) -> str:
    raw = _text(value)
    if not raw or raw in {"0", "00", "000"}:
        return ""
    digits = re.sub(r"[\s.]", "", raw)
    return digits if digits and digits != "0" else ""


def _normal_company_code(value: Any) -> str:
    raw = _text(value)
    if re.fullmatch(r"\d+\.0", raw):
        raw = raw[:-2]
    return raw.strip()


def _normal_marital(value: Any) -> str:
    code = _text(value).strip().upper()
    return {
        "S": "Soltero/a",
        "C": "Casado/a",
        "D": "Divorciado/a",
        "V": "Viudo/a",
        "U": "Unión de hecho",
    }.get(code, "No informado")


def _time_text(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if hasattr(value, "strftime") and not isinstance(value, date):
        try:
            return value.strftime("%H:%M")
        except Exception:
            pass
    raw = _text(value)
    if re.fullmatch(r"\d{1,2}:\d{2}(:\d{2})?", raw):
        return raw[:5]
    return raw[:20]


def _row_to_record(headers: list[str], values: tuple[Any, ...], excel_row: int, epoch: datetime | None) -> dict[str, Any]:
    raw = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
    name = " ".join(part for part in (_text(raw.get("NOMTRA")), _text(raw.get("APETRA"))) if part).strip()
    code = _normal_company_code(raw.get("CODEMP"))
    codtra = _text(raw.get("CODTRA"))
    ci = _normal_ci(raw.get("CEDIDE"))
    return {
        "excel_row": excel_row,
        "company_code": code,
        "legacy_employee_code": codtra,
        "document_number": ci,
        "full_name": name,
        "position": (_text(raw.get("OCUPAC")) or _text(raw.get("PROFES")) or "No informado")[:140],
        "profession": _text(raw.get("PROFES"))[:160],
        "admission_date": _date(raw.get("FECENT"), epoch),
        "termination_date": _date(raw.get("FECSAL"), epoch),
        "termination_present": raw.get("FECSAL") not in (None, ""),
        "birth_date": _date(raw.get("FECNAC"), epoch),
        "salary": _money(raw.get("SALACT")),
        "ips_raw": _text(raw.get("IPSXXX")).upper(),
        "ips_number": _digits(raw.get("NUMIPS")),
        "phone": _text(raw.get("TELEFO"))[:60],
        "address": _text(raw.get("DOMACT"))[:240],
        "nationality": (_text(raw.get("NACION")) or "No informado")[:80],
        "sex": {"M": "Masculino", "F": "Femenino"}.get(_text(raw.get("SEXOXX")).upper(), ""),
        "marital_status": _normal_marital(raw.get("ESTCIV")),
        "birth_place": _text(raw.get("LUGNAC"))[:120],
        "district": _text(raw.get("CIUDAD"))[:100],
        "dependent_children": max(0, min(50, int(_money(raw.get("CANHIJ")) or 0))),
        "schedule": _schedule(raw),
        "section_code": _text(raw.get("CODSEC")),
        "raw_salary_type": _text(raw.get("HORIST")),
        "raw_category": _text(raw.get("CATEGO")),
        "raw_education": _text(raw.get("NIVEST")),
        "raw_ips": _text(raw.get("IPSXXX")),
        "raw_legacy": {
            "CODTRA": codtra,
            "CODEMP": code,
            "CODSEC": _text(raw.get("CODSEC")),
            "CEDIDE": _text(raw.get("CEDIDE")),
            "HORIST": _text(raw.get("HORIST")),
            "HORNOR": _text(raw.get("HORNOR")),
            "CATEGO": _text(raw.get("CATEGO")),
            "PERPRU": _text(raw.get("PERPRU")),
            "REMCON": _text(raw.get("REMCON")),
        },
    }


def _schedule(raw: dict[str, Any]) -> str:
    periods = []
    for start, end, label in (
        ("ENTMAN", "SALMAN", "Mañana"),
        ("ENTTAR", "SALTAR", "Tarde"),
        ("ENTNOC", "SALNOC", "Noche"),
        ("ENTCON", "SALCON", "Complementario"),
    ):
        begin = _time_text(raw.get(start))
        finish = _time_text(raw.get(end))
        if begin or finish:
            periods.append(f"{label}: {begin or '—'}–{finish or '—'}")
    return " · ".join(periods)[:120] if periods else "No informado"


def parse_legacy_employee_workbook(
    content: bytes,
    filename: str,
    companies_by_legacy_code: dict[str, list[dict[str, Any]]],
    *,
    max_rows: int = 15_000,
    today: date | None = None,
) -> dict[str, Any]:
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("La base histórica debe estar en formato .xlsx.")
    if len(content) > 8 * 1024 * 1024:
        raise ValueError("El archivo supera el máximo permitido de 8 MB.")

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(iterator)
        except StopIteration:
            raise ValueError("El libro no contiene datos.")
        headers = [_text(value).strip().upper() for value in raw_headers]
        if len(set(header for header in headers if header)) != len([header for header in headers if header]):
            raise ValueError("El libro contiene encabezados repetidos.")
        missing = sorted(REQUIRED_COLUMNS.difference(headers))
        if missing:
            raise ValueError("Faltan columnas requeridas: " + ", ".join(missing))

        grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
        source_rows = 0
        for excel_row, values in enumerate(iterator, start=2):
            if not any(value not in (None, "") for value in values):
                continue
            source_rows += 1
            if source_rows > max_rows:
                raise ValueError(f"El libro supera el máximo permitido de {max_rows} filas.")
            record = _row_to_record(headers, values, excel_row, workbook.epoch)
            if not record["company_code"]:
                key = ("", f"ROW:{excel_row}")
            elif record["document_number"]:
                key = (record["company_code"], f"CI:{record['document_number']}")
            else:
                key = (record["company_code"], f"COD:{record['legacy_employee_code'] or excel_row}")
            grouped[key].append(record)

        if source_rows == 0:
            raise ValueError("El libro no contiene filas de trabajadores.")

        candidates: list[dict[str, Any]] = []
        excluded: list[dict[str, Any]] = []
        repeated_groups = 0
        no_ci_groups = 0
        invalid_birth_groups = 0
        missing_salary_groups = 0
        invalid_termination_groups = 0
        open_groups = 0
        closed_groups = 0
        group_today = today or date.today()

        for (company_code, identity_key), records in grouped.items():
            records.sort(key=lambda item: item["excel_row"])
            if len(records) > 1:
                repeated_groups += 1
            if identity_key.startswith("COD:"):
                no_ci_groups += 1

            company_matches = companies_by_legacy_code.get(company_code, [])
            if len(company_matches) != 1:
                excluded.append({
                    "company_code": company_code,
                    "name": next((r["full_name"] for r in records if r["full_name"]), ""),
                    "document_number": records[0]["document_number"],
                    "excel_row": records[0]["excel_row"],
                    "reason": "El código patronal no coincide con una única empresa del estudio.",
                })
                continue

            valid_admissions = [record for record in records if record["admission_date"]]
            names = [record for record in records if record["full_name"]]
            if not valid_admissions:
                excluded.append({
                    "company_code": company_code,
                    "company_name": company_matches[0]["name"],
                    "name": names[-1]["full_name"] if names else "",
                    "document_number": records[0]["document_number"],
                    "excel_row": records[0]["excel_row"],
                    "reason": "No tiene una fecha de ingreso válida en ninguna fila del legajo.",
                })
                continue
            if not names:
                excluded.append({
                    "company_code": company_code,
                    "company_name": company_matches[0]["name"],
                    "name": "",
                    "document_number": records[0]["document_number"],
                    "excel_row": records[0]["excel_row"],
                    "reason": "No tiene nombre ni apellido en ninguna fila del legajo.",
                })
                continue

            open_records = [record for record in records if record["termination_date"] is None]
            current_employment = bool(open_records)
            preferred_pool = open_records or records
            preferred = max(
                preferred_pool,
                key=lambda record: (
                    record["admission_date"] or date.min,
                    record["excel_row"],
                ),
            )
            admission = preferred["admission_date"] or max(
                valid_admissions, key=lambda record: (record["admission_date"], record["excel_row"])
            )["admission_date"]
            chosen_name = preferred["full_name"] or names[-1]["full_name"]
            chosen_position = preferred["position"] if preferred["position"] != "No informado" else next(
                (record["position"] for record in reversed(records) if record["position"] != "No informado"),
                "No informado",
            )
            salary_record = preferred if preferred["salary"] is not None else next(
                (record for record in sorted(records, key=lambda r: (r["admission_date"] or date.min, r["excel_row"]), reverse=True)
                 if record["salary"] is not None),
                None,
            )
            salary = salary_record["salary"] if salary_record and salary_record["salary"] is not None else 0
            if salary_record is None:
                missing_salary_groups += 1

            birth_record = next(
                (record for record in reversed(records)
                 if record["birth_date"] and record["birth_date"] < admission and record["birth_date"] <= group_today),
                None,
            )
            birth = birth_record["birth_date"] if birth_record else None
            if any(record["birth_date"] and (
                record["birth_date"] >= (record["admission_date"] or admission)
                or record["birth_date"] > group_today
            ) for record in records):
                invalid_birth_groups += 1

            all_termination_dates = [
                record["termination_date"] for record in records
                if record["termination_date"] and (
                    not record["admission_date"] or record["termination_date"] >= record["admission_date"]
                )
            ]
            termination = None if current_employment else (max(all_termination_dates) if all_termination_dates else None)
            company = company_matches[0]
            if current_employment:
                open_groups += 1
            else:
                closed_groups += 1

            document = preferred["document_number"] or records[0]["document_number"]
            synthetic_document = not bool(document)
            if synthetic_document:
                codtra = preferred["legacy_employee_code"] or records[0]["legacy_employee_code"] or str(records[0]["excel_row"])
                document = f"S/D-{company_code}-{codtra}"[:40]

            ips_value = preferred["ips_raw"]
            ips_contributor = ips_value == "S"
            if ips_value not in {"S", "N"}:
                ips_contributor = False

            source_codtras = sorted({record["legacy_employee_code"] for record in records if record["legacy_employee_code"]})
            notes = [
                f"Importación histórica: {filename}.",
                f"Código(s) histórico(s) CODTRA: {', '.join(source_codtras) or 'No informado'}.",
                f"Filas de origen: {', '.join(str(record['excel_row']) for record in records)}.",
                "El archivo no informa tipo de contrato ni frecuencia de pago; deben verificarse.",
                "La jornada de 8 horas diarias y 48 semanales es un valor provisional del sistema y debe verificarse.",
            ]
            if synthetic_document:
                notes.append("La cédula no estaba informada; se asignó un identificador interno S/D, no una cédula real.")
            if birth is None and any(record["birth_date"] for record in records):
                notes.append("Fecha de nacimiento de origen no utilizada por inconsistencia con la fecha de ingreso.")
            if salary_record is None:
                notes.append("Salario actual no informado en la fuente; se dejó en Gs. 0 para revisión.")
            if invalid_termination:
                invalid_termination_groups += 1
                notes.append("Una o más fechas de egreso de la fuente son inválidas o anteriores al ingreso; no se cargó una fecha de egreso dudosa.")
            if not any(record["admission_date"] for record in open_records) and current_employment:
                notes.append("El registro abierto no indica ingreso; se usa la última fecha de ingreso disponible del legajo.")
            if current_employment:
                employee_status = "Activo"
            else:
                employee_status = "Desvinculado"

            periods = []
            for record in records:
                base_detail = (
                    f"Archivo: {filename}; fila Excel: {record['excel_row']}; "
                    f"CODEMP: {company_code}; CODTRA: {record['legacy_employee_code'] or 'No informado'}; "
                    f"CODSEC: {record['section_code'] or 'No informado'}; "
                    f"Cargo fuente: {record['position']}; salario registrado: "
                    f"{record['salary'] if record['salary'] is not None else 'No informado'}."
                )
                if record["admission_date"]:
                    periods.append({
                        "event_type": "Reingreso histórico" if len(records) > 1 and len(periods) > 0 else "Ingreso histórico",
                        "title": f"Ingreso histórico · {record['legacy_employee_code'] or 'sin CODTRA'}",
                        "detail": base_detail,
                        "effective_date": record["admission_date"].isoformat(),
                        "amount": record["salary"],
                    })
                if record["termination_date"] and (
                    not record["admission_date"] or record["termination_date"] >= record["admission_date"]
                ):
                    periods.append({
                        "event_type": "Egreso histórico",
                        "title": f"Egreso histórico · {record['legacy_employee_code'] or 'sin CODTRA'}",
                        "detail": base_detail,
                        "effective_date": record["termination_date"].isoformat(),
                        "amount": None,
                    })

            schedule = preferred["schedule"]
            if not schedule or schedule == "No informado":
                schedule = next((record["schedule"] for record in reversed(records) if record["schedule"] != "No informado"), "No informado")
            profile = {
                "document_type": "CI" if not synthetic_document else "Otro",
                "sex": preferred["sex"] or "",
                "nationality": preferred["nationality"] or "No informado",
                "marital_status": preferred["marital_status"],
                "birth_place": preferred["birth_place"] or "",
                "department": company.get("department") or "",
                "district": preferred["district"] or "",
                "profession": preferred["profession"] or "",
                "occupation_code": "",
                "position_category": "",
                "work_schedule": schedule,
                "shift": "Diurno",
                "weekly_hours": 48,
                "salary_type": "No especificado",
                "daily_hours": 8,
                "ips_number": preferred["ips_number"] or "",
                "dependent_children": preferred["dependent_children"],
                "indigenous": False,
                "disability": False,
                "rei_status": "Pendiente",
                "reop_status": "Pendiente",
            }
            candidates.append({
                "company_id": company["id"],
                "company_code": company_code,
                "company_name": company["name"],
                "company_status": company["status"],
                "branch_id": company.get("branch_id"),
                "full_name": chosen_name[:180],
                "document_number": document,
                "synthetic_document": synthetic_document,
                "birth_date": birth.isoformat() if birth else None,
                "position": chosen_position[:140],
                "admission_date": admission.isoformat(),
                "termination_date": termination.isoformat() if termination else None,
                "base_salary": salary,
                "ips_contributor": ips_contributor,
                "phone": preferred["phone"],
                "address": preferred["address"],
                "status": employee_status,
                "legacy_employee_code": preferred["legacy_employee_code"],
                "source_codtras": source_codtras,
                "source_rows": [record["excel_row"] for record in records],
                "notes": " ".join(notes),
                "profile": profile,
                "events": periods,
                "has_rehire_history": len(records) > 1,
                "active_in_source": current_employment,
                "warning_invalid_birth": bool(any(record["birth_date"] and (
                    record["birth_date"] >= (record["admission_date"] or admission)
                    or record["birth_date"] > group_today
                ) for record in records)),
            })

        candidates.sort(key=lambda item: (item["company_name"].casefold(), item["full_name"].casefold()))
        excluded.sort(key=lambda item: (item.get("company_name", "").casefold(), item["excel_row"]))
        summary = {
            "source_rows": source_rows,
            "identity_groups": len(grouped),
            "importable": len(candidates),
            "excluded": len(excluded),
            "active": sum(1 for item in candidates if item["active_in_source"]),
            "terminated": sum(1 for item in candidates if not item["active_in_source"]),
            "repeated_identity_groups": repeated_groups,
            "without_cedula": sum(1 for item in candidates if item["synthetic_document"]),
            "invalid_birth_date": sum(1 for item in candidates if item["warning_invalid_birth"]),
            "without_salary": missing_salary_groups,
            "invalid_termination_dates": invalid_termination_groups,
            "historical_events": sum(len(item["events"]) for item in candidates),
            "archived_company_employees": sum(1 for item in candidates if item["company_status"] == "Archivada"),
        }
        return {
            "filename": filename,
            "summary": summary,
            "candidates": candidates,
            "excluded": excluded,
        }
    finally:
        workbook.close()
