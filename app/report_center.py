"""Data-backed report catalog and exports for the labor-management workspace.

The catalog deliberately reuses canonical DIGIT-LABORAL records. It does not
fabricate payroll, IPS, MTESS, or termination data to fill an empty report.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from xml.sax.saxutils import escape as xml_escape
from sqlalchemy import String, cast, func, or_, select
from sqlalchemy.orm import Session

from .models import (
    Aguinaldo,
    CalculationRecord,
    Company,
    Employee,
    EmployeeEvent,
    GeneratedCertificate,
    LaborDeadline,
    LaborOperation,
    MedicalExam,
    OfficialCertificateRecord,
    Payroll,
    PayrollComplianceDetail,
    PayrollLine,
    Vacation,
)


@dataclass(frozen=True)
class ReportSpec:
    key: str
    label: str
    group: str
    description: str
    dataset: str
    variant: str = ""
    note: str = ""




def report_sort_options(spec: ReportSpec) -> list[tuple[str, str]]:
    """Return only sort keys supported by the selected report dataset."""
    options = {
        "company": [("name", "Empresa"), ("code", "Código"), ("city", "Ciudad"), ("status", "Estado")],
        "employee": [("name", "Funcionario"), ("date", "Fecha de ingreso"), ("salary", "Salario base"), ("code", "Código"), ("status", "Estado")],
        "calculation": [("date", "Fecha"), ("type", "Tipo de cálculo"), ("period", "Periodo"), ("amount", "Importe"), ("status", "Estado")],
        "payroll": [("period", "Periodo"), ("amount", "Neto"), ("gross", "Bruto"), ("status", "Estado")],
        "payroll_line": [("period", "Periodo"), ("employee", "Funcionario"), ("amount", "Neto"), ("gross", "Bruto"), ("status", "Estado")],
        "family_allowance": [("period", "Periodo"), ("employee", "Funcionario"), ("amount", "Importe"), ("status", "Estado")],
        "aguinaldo": [("period", "Año"), ("employee", "Funcionario"), ("amount", "Importe"), ("status", "Estado"), ("date", "Fecha de registro")],
        "vacation": [("period", "Año"), ("employee", "Funcionario"), ("status", "Estado"), ("date", "Fecha de inicio")],
        "operation": [("date", "Fecha"), ("name", "Funcionario"), ("amount", "Importe"), ("status", "Estado")],
        "medical_exam": [("date", "Fecha de examen"), ("employee", "Funcionario"), ("status", "Resultado"), ("amount", "Vencimiento")],
        "certificate": [("date", "Fecha"), ("name", "Funcionario"), ("status", "Estado")],
        "employee_event": [("date", "Fecha"), ("name", "Funcionario"), ("status", "Tipo de evento")],
        "deadline": [("date", "Vencimiento"), ("name", "Empresa"), ("status", "Estado")],
    }
    return options.get(spec.dataset, [("name", "Nombre"), ("date", "Fecha"), ("status", "Estado")])



_REPORT_LIST = [
    ("liquidacion-salario", "Liquidación de Salarios", "Liquidaciones y remuneraciones", "Recibos detallados de cálculos salariales guardados, con historial, impresión y PDF.", "calculation", "salary", ""),
    ("liq-bonificacion-flia", "Liq. Bonificación Familiar", "Liquidaciones y remuneraciones", "Detalle de bonificación familiar cuando se encuentre registrado en la nómina.", "family_allowance", "", "Solo refleja importes registrados; no determina por sí solo el derecho legal a la bonificación."),
    ("liquidacion-haberes", "Liquidación Haberes", "Liquidaciones y remuneraciones", "Liquidaciones de haberes generadas y guardadas para revisión.", "calculation", "settlement", "El reporte histórico no sustituye la revisión profesional de una liquidación final."),
    ("listado-pre-aviso", "Listado Pre-Aviso", "Liquidaciones y remuneraciones", "Cálculos de preaviso registrados.", "calculation", "notice", "La procedencia y el plazo deben verificarse según el caso concreto."),
    ("listado-empl-fecha", "Listado Empl. Fecha", "Legajos y empresas", "Funcionarios ordenados por fecha de ingreso.", "employee", "admission_date", ""),
    ("personales-despedidos", "Personales Despedidos", "Legajos y empresas", "Funcionarios con fecha de egreso o estado que indique baja/despido.", "employee", "terminated", "Solo identifica datos registrados en el legajo; no determina la causal jurídica."),
    ("informes-salario", "Informes Salario", "Liquidaciones y remuneraciones", "Historial de resultados de cálculos de salario.", "calculation", "salary", ""),
    ("seguro-social", "Seguro Social", "IPS y cumplimiento", "Nóminas y datos de aportes al seguro social registrados.", "payroll_line", "ips", "Los valores son los registrados en DIGIT-LABORAL; confirmar con los documentos oficiales de IPS."),
    ("planilla-de-pago", "Planilla de Pago", "Liquidaciones y remuneraciones", "Detalle de conceptos, descuentos y neto de nómina.", "payroll_line", "", ""),
    ("liquidacion-aguinaldo", "Liquidación Aguinaldo", "Vacaciones y aguinaldo", "Cálculos de aguinaldo guardados para revisión.", "calculation", "aguinaldo", "No reemplaza la validación de remuneraciones computables del periodo."),
    ("planilla-de-horas", "Planilla de Horas", "Liquidaciones y remuneraciones", "Horas trabajadas y conceptos horarios asociados a nóminas registradas.", "payroll_line", "hours", ""),
    ("listado-vacacion", "Listado Vacación", "Vacaciones y aguinaldo", "Saldos y periodos de vacaciones registrados.", "vacation", "", "El saldo mostrado se obtiene de los días registrados; debe revisarse antes de liquidar."),
    ("listado-empresa", "Listado Empresa", "Legajos y empresas", "Directorio patronal del estudio.", "company", "", ""),
    ("listado-empleado", "Listado Empleado", "Legajos y empresas", "Listado de funcionarios y datos laborales principales.", "employee", "", ""),
    ("ficha-empleado", "Ficha Empleado", "Legajos y empresas", "Ficha de datos del funcionario seleccionado.", "employee", "profile", "Elegí un funcionario para consultar su ficha completa."),
    ("bonificacion-familiar", "Bonificación Familiar", "Liquidaciones y remuneraciones", "Bonificación familiar ingresada en los detalles complementarios de nómina.", "family_allowance", "", "No infiere dependientes ni elegibilidad si esos datos no fueron registrados."),
    ("listado-anticipo", "Listado Anticipo", "Liquidaciones y remuneraciones", "Anticipos de salario registrados como operaciones laborales.", "operation", "advance", ""),
    ("ips-pre-impreso-1", "I.P.S. Pre-Impreso-1", "IPS y cumplimiento", "Vista preliminar de conceptos de nómina asociados a IPS.", "payroll_line", "ips", "Formato interno de preparación; no es un formulario oficial ni una presentación a IPS."),
    ("ips-pre-impreso-2", "I.P.S. Pre-Impreso-2", "IPS y cumplimiento", "Segunda vista de preparación de conceptos de IPS y nómina.", "payroll_line", "ips", "Formato interno de preparación; no es un formulario oficial ni una presentación a IPS."),
    ("comunicacion", "Comunicación", "Trámites y seguimiento", "Comunicaciones laborales registradas en operaciones y eventos.", "operation", "communication", "Los borradores no acreditan notificación efectiva ni recepción."),
    ("salario-pre-impreso", "Salario Pre-Impreso", "Liquidaciones y remuneraciones", "Vista de impresión previa de los importes de nómina.", "payroll_line", "", "El documento debe revisarse y aprobarse antes de utilizarlo como recibo."),
    ("planilla-de-aguinaldo", "Planilla de Aguinaldo", "Vacaciones y aguinaldo", "Resumen anual de aguinaldos registrados.", "aguinaldo", "", ""),
    ("planilla-pago-seccion", "Planilla Pago/Sección", "Liquidaciones y remuneraciones", "Detalle de nómina con sección organizativa del funcionario.", "payroll_line", "section", ""),
    ("listado-notificacion", "Listado Notificación", "Trámites y seguimiento", "Recordatorios y vencimientos laborales registrados.", "deadline", "", "El reporte muestra vencimientos registrados; no envía ni confirma notificaciones por sí solo."),
    ("p-sueldos-jornales", "P. Sueldos y Jornales", "Registros laborales anuales", "Resumen de conceptos de remuneración y jornales registrados en nómina.", "payroll_line", "", ""),
    ("p-empleados-obreros", "P. Empleados y Obreros", "Registros laborales anuales", "Nómina de empleados y obreros registrados.", "employee", "", ""),
    ("p-reg-gral-personas-ocupadas", "P.R. Gral. Personas Ocupadas", "Registros laborales anuales", "Listado de personas ocupadas según los legajos existentes.", "employee", "", "No realiza una clasificación ocupacional oficial si el perfil no está completado."),
    ("p-sueldos-jornales-mensual", "P. Sueldos y Jornales Mensual", "Registros laborales anuales", "Detalle mensual de conceptos de nómina.", "payroll_line", "", ""),
    ("p-vacaciones-anuales", "P. Vacaciones Anuales", "Vacaciones y aguinaldo", "Control anual de días de vacaciones y periodos registrados.", "vacation", "", ""),
    ("liquidacion-vacacion", "Liquidación Vacación", "Vacaciones y aguinaldo", "Cálculos de vacaciones guardados para revisión.", "calculation", "vacation", "No genera automáticamente importes debidos a partir de datos incompletos."),
    ("planilla-de-horas-1", "Planilla de Horas-1", "Liquidaciones y remuneraciones", "Segunda vista de operaciones de horas, turnos y recargos.", "operation", "hours", ""),
    ("listado-historico", "Listado Histórico", "Trámites y seguimiento", "Eventos laborales históricos vinculados al legajo.", "employee_event", "", ""),
    ("datos-de-la-empresa", "Datos de la Empresa", "Legajos y empresas", "Datos patronales y de contacto registrados en el sistema.", "company", "extended", ""),
    ("listado-anual-reg-empleados-obreros", "L. Anual Reg. Empl. Obr.", "Registros laborales anuales", "Base de empleados y obreros para revisión del registro anual.", "employee", "annual", "Reporte preparatorio; no equivale a una declaración oficial presentada."),
    ("listado-examen-medico", "Listado Examen Médico", "Salud ocupacional", "Exámenes de salud ocupacional, resultados y vencimientos.", "medical_exam", "", "Limitar el acceso y la circulación de esta información a usuarios autorizados."),
]

REPORTS: dict[str, ReportSpec] = {
    key: ReportSpec(key, label, group, description, dataset, variant, note)
    for key, label, group, description, dataset, variant, note in _REPORT_LIST
}

REPORT_GROUPS: list[dict[str, Any]] = []
for _group in dict.fromkeys(spec.group for spec in REPORTS.values()):
    REPORT_GROUPS.append({
        "title": _group,
        "reports": [spec for spec in REPORTS.values() if spec.group == _group],
    })


def get_report(report_key: str) -> ReportSpec | None:
    return REPORTS.get(report_key)


def report_columns(spec: ReportSpec) -> list[str]:
    columns = {
        "company": ["Código", "Empresa", "Nombre comercial", "RUC", "Rubro", "Representante legal", "Contacto", "Teléfono", "Correo", "Dirección", "Ciudad", "N.º patronal IPS", "N.º patronal MTESS", "Estado"],
        "employee": ["Código funcionario", "Funcionario", "Empresa", "Documento", "Cargo", "Ingreso", "Egreso", "Salario base (Gs.)", "Tipo de contrato", "Frecuencia de pago", "Aportante IPS", "Sección", "Correo", "Teléfono", "Estado"],
        "payroll": ["Periodo", "Empresa", "Estado", "Bruto total (Gs.)", "IPS trabajador (Gs.)", "Descuentos (Gs.)", "Neto (Gs.)", "Creado", "Cerrado"],
        "payroll_line": ["Periodo", "Empresa", "Funcionario", "Sección", "Salario base (Gs.)", "Horas extras (Gs.)", "Comisiones (Gs.)", "Bonificaciones (Gs.)", "Otros ingresos (Gs.)", "Anticipos (Gs.)", "IPS trabajador (Gs.)", "Descuentos (Gs.)", "Bruto (Gs.)", "Neto (Gs.)", "Días trabajados", "Horas trabajadas", "Bonificación familiar (Gs.)", "IPS empleador (Gs.)"],
        "calculation": ["Fecha", "Empresa", "Funcionario", "Tipo de cálculo", "Periodo", "Importe (Gs.)", "Estado", "Origen", "Registrado por"],
        "aguinaldo": ["Año", "Empresa", "Funcionario", "Remuneraciones (Gs.)", "Aguinaldo calculado (Gs.)", "Estado", "Registrado"],
        "vacation": ["Año de periodo", "Empresa", "Funcionario", "Días adquiridos", "Días utilizados", "Días pendientes (referencial)", "Desde", "Hasta", "Estado"],
        "medical_exam": ["Código de examen", "Empresa", "Funcionario", "Tipo", "Fecha de examen", "Resultado", "Vencimiento", "Estado de vigencia", "Observaciones"],
        "operation": ["Fecha efectiva", "Empresa", "Funcionario", "Tipo de operación", "Código", "Detalle", "Cantidad", "Importe (Gs.)", "Estado", "Observaciones"],
        "certificate": ["Fecha", "Empresa", "Funcionario", "Documento", "Tipo", "Estado", "Creado por"],
        "employee_event": ["Fecha efectiva", "Empresa", "Funcionario", "Tipo de evento", "Título", "Detalle", "Importe (Gs.)", "Registrado por"],
        "deadline": ["Vencimiento", "Empresa", "Funcionario", "Tipo", "Descripción", "Prioridad", "Estado", "Observaciones"],
        "family_allowance": ["Periodo", "Empresa", "Funcionario", "Bonificación familiar (Gs.)", "Días trabajados", "Pago", "Estado de nómina"],
    }
    if spec.dataset == "employee" and spec.variant == "profile":
        return columns["employee"] + ["Dirección", "Fecha de nacimiento"]
    return columns.get(spec.dataset, [])




def _calculation_label(value: str) -> str:
    key = (value or "").removeprefix("suite_")
    legacy = {"salary": "Salario neto", "hours": "Horas", "aguinaldo": "Aguinaldo", "vacation": "Vacaciones", "notice": "Preaviso", "settlement": "Liquidación de haberes"}
    if key in legacy:
        return legacy[key]
    try:
        from .labor_calculator import TITLES
        return TITLES.get(key, key.replace("_", " ").capitalize())
    except Exception:
        return key.replace("_", " ").capitalize()


def _clean_search(value: str | None) -> str:
    return (value or "").strip().replace("%", "").replace("_", "")[:120]


def _date(value: date | datetime | None) -> str:
    return value.isoformat() if value else ""


def _bool(value: bool | None) -> str:
    return "Sí" if value else "No"


def _company_query(db: Session, company_ids: list[int], company_id: int | None, q: str):
    query = select(Company).where(Company.id.in_(company_ids))
    count_query = select(func.count(Company.id)).where(Company.id.in_(company_ids))
    if company_id is not None:
        query = query.where(Company.id == company_id)
        count_query = count_query.where(Company.id == company_id)
    if q:
        term = f"%{q}%"
        clause = or_(Company.legal_name.ilike(term), Company.trade_name.ilike(term), Company.ruc.ilike(term), Company.city.ilike(term), cast(Company.company_code, String).ilike(term), Company.work_sector.ilike(term))
        query = query.where(clause)
        count_query = count_query.where(clause)
    return query, count_query


def _employee_query(company_ids: list[int], company_id: int | None, employee_id: int | None, q: str, spec: ReportSpec, year: int | None):
    query = select(Employee).where(Employee.company_id.in_(company_ids))
    count_query = select(func.count(Employee.id)).where(Employee.company_id.in_(company_ids))
    if company_id is not None:
        query, count_query = query.where(Employee.company_id == company_id), count_query.where(Employee.company_id == company_id)
    if employee_id is not None:
        query, count_query = query.where(Employee.id == employee_id), count_query.where(Employee.id == employee_id)
    if spec.variant == "terminated":
        clause = or_(Employee.termination_date.is_not(None), Employee.status.ilike("%baja%"), Employee.status.ilike("%desped%"), Employee.status.ilike("%egreso%"))
        query, count_query = query.where(clause), count_query.where(clause)
    if year is not None:
        clause = func.extract("year", Employee.admission_date) == year
        query, count_query = query.where(clause), count_query.where(clause)
    if q:
        term = f"%{q}%"
        clause = or_(Employee.full_name.ilike(term), Employee.document_number.ilike(term), Employee.position.ilike(term), Employee.status.ilike(term), cast(Employee.employee_code, String).ilike(term))
        query, count_query = query.where(clause), count_query.where(clause)
    return query, count_query


def _payroll_query(company_ids: list[int], company_id: int | None, q: str, status_filter: str, year: int | None, period: str | None):
    query = select(Payroll).where(Payroll.company_id.in_(company_ids))
    count_query = select(func.count(Payroll.id)).where(Payroll.company_id.in_(company_ids))
    if company_id is not None:
        query, count_query = query.where(Payroll.company_id == company_id), count_query.where(Payroll.company_id == company_id)
    if q:
        term = f"%{q}%"
        clause = or_(Payroll.period.ilike(term), Payroll.status.ilike(term), Payroll.notes.ilike(term))
        query, count_query = query.where(clause), count_query.where(clause)
    if status_filter:
        query, count_query = query.where(Payroll.status == status_filter), count_query.where(Payroll.status == status_filter)
    if year is not None:
        clause = Payroll.period.startswith(str(year))
        query, count_query = query.where(clause), count_query.where(clause)
    if period:
        clause = Payroll.period == period
        query, count_query = query.where(clause), count_query.where(clause)
    return query, count_query


def _payroll_line_query(company_ids: list[int], company_id: int | None, employee_id: int | None, q: str, status_filter: str, year: int | None, period: str | None, spec: ReportSpec):
    from .models import CompanySection  # kept local to keep the common model imports concise

    query = select(PayrollLine, Payroll, Employee, Company, PayrollComplianceDetail).join(
        Payroll, PayrollLine.payroll_id == Payroll.id
    ).join(Employee, PayrollLine.employee_id == Employee.id).join(
        Company, Payroll.company_id == Company.id
    ).outerjoin(PayrollComplianceDetail, PayrollComplianceDetail.payroll_line_id == PayrollLine.id).where(
        Payroll.company_id.in_(company_ids)
    )
    count_query = select(func.count(PayrollLine.id)).select_from(PayrollLine).join(
        Payroll, PayrollLine.payroll_id == Payroll.id
    ).join(Employee, PayrollLine.employee_id == Employee.id).join(
        Company, Payroll.company_id == Company.id
    ).outerjoin(PayrollComplianceDetail, PayrollComplianceDetail.payroll_line_id == PayrollLine.id).where(
        Payroll.company_id.in_(company_ids)
    )
    if company_id is not None:
        clause = Payroll.company_id == company_id
        query, count_query = query.where(clause), count_query.where(clause)
    if employee_id is not None:
        clause = Employee.id == employee_id
        query, count_query = query.where(clause), count_query.where(clause)
    if q:
        term = f"%{q}%"
        clause = or_(Employee.full_name.ilike(term), Employee.document_number.ilike(term), Company.legal_name.ilike(term), Payroll.period.ilike(term))
        query, count_query = query.where(clause), count_query.where(clause)
    if status_filter:
        clause = Payroll.status == status_filter
        query, count_query = query.where(clause), count_query.where(clause)
    if year is not None:
        clause = Payroll.period.startswith(str(year))
        query, count_query = query.where(clause), count_query.where(clause)
    if period:
        clause = Payroll.period == period
        query, count_query = query.where(clause), count_query.where(clause)
    if spec.dataset == "family_allowance":
        query = query.where(PayrollComplianceDetail.family_allowance > 0)
        count_query = count_query.where(PayrollComplianceDetail.family_allowance > 0)
    if spec.variant == "hours":
        # Hours are persisted in PayrollComplianceDetail, not inferred from gross pay.
        query = query.where(PayrollComplianceDetail.hours_worked > 0)
        count_query = count_query.where(PayrollComplianceDetail.hours_worked > 0)
    if spec.variant == "ips":
        query = query.where(or_(PayrollLine.ips_employee > 0, PayrollComplianceDetail.employer_ips > 0))
        count_query = count_query.where(or_(PayrollLine.ips_employee > 0, PayrollComplianceDetail.employer_ips > 0))
    if spec.variant == "section":
        query = query.join(Employee.section)
        count_query = count_query.join(Employee.section)
    return query, count_query


def _calculation_query(company_ids: list[int], company_id: int | None, employee_id: int | None, q: str, status_filter: str, year: int | None, spec: ReportSpec):
    query = select(CalculationRecord, Company, Employee).join(
        Company, CalculationRecord.company_id == Company.id
    ).outerjoin(Employee, CalculationRecord.employee_id == Employee.id).where(
        CalculationRecord.company_id.in_(company_ids)
    )
    count_query = select(func.count(CalculationRecord.id)).select_from(CalculationRecord).where(
        CalculationRecord.company_id.in_(company_ids)
    )
    if company_id is not None:
        clause = CalculationRecord.company_id == company_id
        query, count_query = query.where(clause), count_query.where(clause)
    if employee_id is not None:
        clause = CalculationRecord.employee_id == employee_id
        query, count_query = query.where(clause), count_query.where(clause)
    if spec.variant:
        # Both legacy calculators and the professional suite persist the same
        # reportable concepts under different, documented type prefixes.
        clause = or_(
            CalculationRecord.calculation_type == spec.variant,
            CalculationRecord.calculation_type == f"suite_{spec.variant}",
        )
        query, count_query = query.where(clause), count_query.where(clause)
    if q:
        term = f"%{q}%"
        clause = or_(Company.legal_name.ilike(term), Employee.full_name.ilike(term), CalculationRecord.calculation_type.ilike(term), CalculationRecord.reference_period.ilike(term), CalculationRecord.notes.ilike(term))
        query, count_query = query.where(clause), count_query.where(clause)
    if status_filter:
        clause = CalculationRecord.status == status_filter
        query, count_query = query.where(clause), count_query.where(clause)
    if year is not None:
        clause = func.extract("year", CalculationRecord.created_at) == year
        query, count_query = query.where(clause), count_query.where(clause)
    return query, count_query


def _related_query(dataset: str, company_ids: list[int], company_id: int | None, employee_id: int | None, q: str, status_filter: str, year: int | None, period: str | None):
    if dataset == "aguinaldo":
        query = select(Aguinaldo, Employee, Company).join(Employee, Aguinaldo.employee_id == Employee.id).join(Company, Employee.company_id == Company.id).where(Employee.company_id.in_(company_ids))
        count_query = select(func.count(Aguinaldo.id)).select_from(Aguinaldo).join(Employee, Aguinaldo.employee_id == Employee.id).where(Employee.company_id.in_(company_ids))
        date_field = Aguinaldo.year
        if year is not None:
            query, count_query = query.where(date_field == year), count_query.where(date_field == year)
        if period:
            try:
                yy = int(period[:4])
                query, count_query = query.where(date_field == yy), count_query.where(date_field == yy)
            except ValueError:
                pass
        if q:
            term = f"%{q}%"
            clause = or_(Employee.full_name.ilike(term), Employee.document_number.ilike(term), Company.legal_name.ilike(term), Aguinaldo.status.ilike(term))
            query, count_query = query.where(clause), count_query.where(clause)
        if status_filter:
            query, count_query = query.where(Aguinaldo.status == status_filter), count_query.where(Aguinaldo.status == status_filter)
        if company_id is not None:
            query, count_query = query.where(Company.id == company_id), count_query.where(Company.id == company_id)
        if employee_id is not None:
            query, count_query = query.where(Employee.id == employee_id), count_query.where(Employee.id == employee_id)
        return query, count_query
    if dataset == "vacation":
        query = select(Vacation, Employee, Company).join(Employee, Vacation.employee_id == Employee.id).join(Company, Employee.company_id == Company.id).where(Employee.company_id.in_(company_ids))
        count_query = select(func.count(Vacation.id)).select_from(Vacation).join(Employee, Vacation.employee_id == Employee.id).where(Employee.company_id.in_(company_ids))
        if year is not None:
            query, count_query = query.where(Vacation.period_year == year), count_query.where(Vacation.period_year == year)
        if period:
            try:
                yy = int(period[:4])
                query, count_query = query.where(Vacation.period_year == yy), count_query.where(Vacation.period_year == yy)
            except ValueError:
                pass
        if q:
            term = f"%{q}%"
            clause = or_(Employee.full_name.ilike(term), Employee.document_number.ilike(term), Company.legal_name.ilike(term), Vacation.status.ilike(term))
            query, count_query = query.where(clause), count_query.where(clause)
        if status_filter:
            query, count_query = query.where(Vacation.status == status_filter), count_query.where(Vacation.status == status_filter)
        if company_id is not None:
            query, count_query = query.where(Company.id == company_id), count_query.where(Company.id == company_id)
        if employee_id is not None:
            query, count_query = query.where(Employee.id == employee_id), count_query.where(Employee.id == employee_id)
        return query, count_query
    if dataset == "medical_exam":
        query = select(MedicalExam, Employee, Company).join(Employee, MedicalExam.employee_id == Employee.id).join(Company, MedicalExam.company_id == Company.id).where(MedicalExam.company_id.in_(company_ids))
        count_query = select(func.count(MedicalExam.id)).where(MedicalExam.company_id.in_(company_ids))
        if year is not None:
            clause = func.extract("year", MedicalExam.exam_date) == year
            query, count_query = query.where(clause), count_query.where(clause)
        if q:
            term = f"%{q}%"
            clause = or_(MedicalExam.exam_code.ilike(term), MedicalExam.exam_type.ilike(term), MedicalExam.result.ilike(term), Employee.full_name.ilike(term), Company.legal_name.ilike(term))
            query, count_query = query.where(clause), count_query.where(clause)
        if status_filter:
            clause = MedicalExam.result == status_filter
            query, count_query = query.where(clause), count_query.where(clause)
        if company_id is not None:
            clause = MedicalExam.company_id == company_id
            query, count_query = query.where(clause), count_query.where(clause)
        if employee_id is not None:
            clause = Employee.id == employee_id
            query, count_query = query.where(clause), count_query.where(clause)
        return query, count_query
    if dataset == "operation":
        query = select(LaborOperation, Employee, Company).join(Employee, LaborOperation.employee_id == Employee.id).join(Company, LaborOperation.company_id == Company.id).where(LaborOperation.company_id.in_(company_ids))
        count_query = select(func.count(LaborOperation.id)).where(LaborOperation.company_id.in_(company_ids))
        if year is not None:
            clause = func.extract("year", LaborOperation.effective_date) == year
            query, count_query = query.where(clause), count_query.where(clause)
        if period:
            try:
                period_year, period_month = (int(part) for part in period.split("-", 1))
                start_date = date(period_year, period_month, 1)
                end_date = date(period_year + (period_month == 12), 1 if period_month == 12 else period_month + 1, 1)
                clause = (LaborOperation.effective_date >= start_date) & (LaborOperation.effective_date < end_date)
                query, count_query = query.where(clause), count_query.where(clause)
            except (ValueError, TypeError):
                pass
        if q:
            term = f"%{q}%"
            clause = or_(LaborOperation.title.ilike(term), LaborOperation.operation_type.ilike(term), LaborOperation.reference_code.ilike(term), Employee.full_name.ilike(term), Company.legal_name.ilike(term))
            query, count_query = query.where(clause), count_query.where(clause)
        if status_filter:
            query, count_query = query.where(LaborOperation.status == status_filter), count_query.where(LaborOperation.status == status_filter)
        if company_id is not None:
            query, count_query = query.where(Company.id == company_id), count_query.where(Company.id == company_id)
        if employee_id is not None:
            query, count_query = query.where(Employee.id == employee_id), count_query.where(Employee.id == employee_id)
        return query, count_query
    if dataset == "certificate":
        query = select(GeneratedCertificate, Company, Employee).join(Company, GeneratedCertificate.company_id == Company.id).outerjoin(Employee, GeneratedCertificate.employee_id == Employee.id).where(GeneratedCertificate.company_id.in_(company_ids))
        count_query = select(func.count(GeneratedCertificate.id)).where(GeneratedCertificate.company_id.in_(company_ids))
        if q:
            term = f"%{q}%"
            clause = or_(GeneratedCertificate.title.ilike(term), GeneratedCertificate.document_type.ilike(term), GeneratedCertificate.employee_name_snapshot.ilike(term), Company.legal_name.ilike(term))
            query, count_query = query.where(clause), count_query.where(clause)
        if status_filter:
            query, count_query = query.where(GeneratedCertificate.status == status_filter), count_query.where(GeneratedCertificate.status == status_filter)
        if company_id is not None:
            query, count_query = query.where(Company.id == company_id), count_query.where(Company.id == company_id)
        if employee_id is not None:
            query, count_query = query.where(GeneratedCertificate.employee_id == employee_id), count_query.where(GeneratedCertificate.employee_id == employee_id)
        if year is not None:
            clause = func.extract("year", GeneratedCertificate.issue_date) == year
            query, count_query = query.where(clause), count_query.where(clause)
        return query, count_query
    if dataset == "employee_event":
        query = select(EmployeeEvent, Employee, Company).join(Employee, EmployeeEvent.employee_id == Employee.id).join(Company, Employee.company_id == Company.id).where(Employee.company_id.in_(company_ids))
        count_query = select(func.count(EmployeeEvent.id)).select_from(EmployeeEvent).join(Employee, EmployeeEvent.employee_id == Employee.id).where(Employee.company_id.in_(company_ids))
        if year is not None:
            clause = func.extract("year", EmployeeEvent.effective_date) == year
            query, count_query = query.where(clause), count_query.where(clause)
        if q:
            term = f"%{q}%"
            clause = or_(EmployeeEvent.title.ilike(term), EmployeeEvent.event_type.ilike(term), Employee.full_name.ilike(term), Company.legal_name.ilike(term), EmployeeEvent.detail.ilike(term))
            query, count_query = query.where(clause), count_query.where(clause)
        if status_filter:
            clause = EmployeeEvent.event_type == status_filter
            query, count_query = query.where(clause), count_query.where(clause)
        if company_id is not None:
            clause = Company.id == company_id
            query, count_query = query.where(clause), count_query.where(clause)
        if employee_id is not None:
            clause = Employee.id == employee_id
            query, count_query = query.where(clause), count_query.where(clause)
        return query, count_query
    if dataset == "deadline":
        query = select(LaborDeadline, Company, Employee).join(Company, LaborDeadline.company_id == Company.id).outerjoin(Employee, LaborDeadline.employee_id == Employee.id).where(LaborDeadline.company_id.in_(company_ids))
        count_query = select(func.count(LaborDeadline.id)).where(LaborDeadline.company_id.in_(company_ids))
        if year is not None:
            clause = func.extract("year", LaborDeadline.due_date) == year
            query, count_query = query.where(clause), count_query.where(clause)
        if q:
            term = f"%{q}%"
            clause = or_(LaborDeadline.title.ilike(term), LaborDeadline.deadline_type.ilike(term), Company.legal_name.ilike(term), Employee.full_name.ilike(term))
            query, count_query = query.where(clause), count_query.where(clause)
        if status_filter:
            clause = LaborDeadline.status == status_filter
            query, count_query = query.where(clause), count_query.where(clause)
        if company_id is not None:
            clause = Company.id == company_id
            query, count_query = query.where(clause), count_query.where(clause)
        if employee_id is not None:
            clause = LaborDeadline.employee_id == employee_id
            query, count_query = query.where(clause), count_query.where(clause)
        return query, count_query
    raise ValueError(f"Fuente de reporte no reconocida: {dataset}")


def _row(spec: ReportSpec, value: Any) -> dict[str, Any]:
    dataset = spec.dataset
    if dataset == "company":
        item = value
        result = [item.company_code or "", item.legal_name, item.trade_name, item.ruc, item.work_sector, item.legal_representative, item.contact_name or item.responsible_name, item.phone, item.email, item.address, item.city, item.ips_employer_number, item.mtess_employer_number, item.status]
        if spec.variant == "extended":
            result[4] = item.work_sector or item.work_location
        return dict(zip(report_columns(spec), result))
    if dataset == "employee":
        item = value
        result = [item.employee_code or "", item.full_name, item.company.legal_name, item.document_number, item.position, _date(item.admission_date), _date(item.termination_date), item.base_salary, item.contract_type, item.payment_frequency, _bool(item.ips_contributor), item.section.name if item.section else "", item.email, item.phone, item.status]
        if spec.variant == "profile":
            result += [item.address, _date(item.birth_date), item.contract_type]
        return dict(zip(report_columns(spec), result))
    if dataset == "payroll":
        item = value
        return dict(zip(report_columns(spec), [item.period, item.company.legal_name, item.status, item.total_gross, item.total_ips_employee, item.total_discounts, item.total_net, _date(item.created_at), _date(item.closed_at)]))
    if dataset == "calculation":
        item, company, employee = value
        name = employee.full_name if employee else ""
        return dict(zip(report_columns(spec), [_date(item.created_at), company.legal_name, name, _calculation_label(item.calculation_type), item.reference_period, item.amount, item.status, item.source, item.created_by]))
    if dataset == "payroll_line":
        line, payroll, employee, company, detail = value
        values = [payroll.period, company.legal_name, employee.full_name, employee.section.name if employee.section else "", line.base_salary, line.overtime, line.commissions, line.bonuses, line.other_income, line.advances, line.ips_employee, line.total_discounts, line.gross, line.net, detail.days_worked if detail else "", str(detail.hours_worked) if detail else "", detail.family_allowance if detail else 0, detail.employer_ips if detail else 0]
        if spec.dataset == "family_allowance":
            values = [payroll.period, company.legal_name, employee.full_name, detail.family_allowance if detail else 0, detail.days_worked if detail else "", _date(payroll.closed_at), payroll.status]
        return dict(zip(report_columns(spec), values))
    if dataset == "family_allowance":
        line, payroll, employee, company, detail = value
        return dict(zip(report_columns(spec), [payroll.period, company.legal_name, employee.full_name, detail.family_allowance if detail else 0, detail.days_worked if detail else "", _date(payroll.closed_at), payroll.status]))
    if dataset == "aguinaldo":
        item, employee, company = value
        return dict(zip(report_columns(spec), [item.year, company.legal_name, employee.full_name, item.total_remunerations, item.calculated_amount, item.status, _date(item.created_at)]))
    if dataset == "vacation":
        item, employee, company = value
        pending = max(0, int(item.entitled_days or 0) - int(item.used_days or 0))
        return dict(zip(report_columns(spec), [item.period_year, company.legal_name, employee.full_name, item.entitled_days, item.used_days, pending, _date(item.start_date), _date(item.end_date), item.status]))
    if dataset == "medical_exam":
        item, employee, company = value
        state = "Vencido" if item.expires_on < date.today() else ("Por vencer" if (item.expires_on - date.today()).days <= 30 else "Vigente")
        return dict(zip(report_columns(spec), [item.exam_code, company.legal_name, employee.full_name, item.exam_type, _date(item.exam_date), item.result, _date(item.expires_on), state, item.observations]))
    if dataset == "operation":
        item, employee, company = value
        return dict(zip(report_columns(spec), [_date(item.effective_date), company.legal_name, employee.full_name, item.operation_type, item.reference_code, item.title, str(item.quantity or ""), item.amount if item.amount is not None else "", item.status, item.notes]))
    if dataset == "certificate":
        item, company, employee = value
        return dict(zip(report_columns(spec), [_date(item.issue_date), company.legal_name, employee.full_name if employee else item.employee_name_snapshot, item.title, item.document_type, item.status, item.created_by]))
    if dataset == "employee_event":
        item, employee, company = value
        return dict(zip(report_columns(spec), [_date(item.effective_date), company.legal_name, employee.full_name, item.event_type, item.title, item.detail, item.amount if item.amount is not None else "", item.created_by]))
    if dataset == "deadline":
        item, company, employee = value
        return dict(zip(report_columns(spec), [_date(item.due_date), company.legal_name, employee.full_name if employee else "", item.deadline_type, item.title, item.priority, item.status, item.notes]))
    return {}


def fetch_report(
    db: Session,
    spec: ReportSpec,
    company_ids: list[int],
    company_id: int | None = None,
    employee_id: int | None = None,
    q: str | None = None,
    status_filter: str | None = None,
    year: int | None = None,
    period: str | None = None,
    page: int = 1,
    page_size: int = 25,
    sort: str | None = None,
) -> dict[str, Any]:
    """Return a tenant-scoped, database-paginated report result."""
    q = _clean_search(q)
    status_filter = (status_filter or "").strip()[:40]
    period = (period or "").strip()[:20]
    page = max(1, min(int(page or 1), 1_000_000))
    page_size = max(1, min(int(page_size or 25), 100))
    if not company_ids:
        return {"records": [], "total": 0, "page": page, "page_size": page_size, "pages": 0, "columns": report_columns(spec)}
    if company_id is not None and company_id not in company_ids:
        return {"records": [], "total": 0, "page": page, "page_size": page_size, "pages": 0, "columns": report_columns(spec)}
    if employee_id is not None:
        employee = db.get(Employee, employee_id)
        if not employee or employee.company_id not in company_ids or (company_id is not None and employee.company_id != company_id):
            return {"records": [], "total": 0, "page": page, "page_size": page_size, "pages": 0, "columns": report_columns(spec)}
    if spec.dataset == "company":
        query, count_query = _company_query(db, company_ids, company_id, q)
        if status_filter:
            query = query.where(Company.status == status_filter)
        sort_map = {"name": Company.legal_name, "code": Company.company_code, "city": Company.city, "status": Company.status}
        order_column = sort_map.get(sort or "name", Company.legal_name)
        query = query.order_by(order_column.asc(), Company.id.asc())
        if spec.variant == "extended":
            pass
    elif spec.dataset == "employee":
        query, count_query = _employee_query(company_ids, company_id, employee_id, q, spec, year)
        if status_filter:
            query = query.where(Employee.status == status_filter)
        sort_map = {"name": Employee.full_name, "admission_date": Employee.admission_date, "date": Employee.admission_date, "salary": Employee.base_salary, "amount": Employee.base_salary, "code": Employee.employee_code, "status": Employee.status}
        order_column = sort_map.get(sort or ("admission_date" if spec.variant == "admission_date" else "name"), Employee.full_name)
        query = query.order_by(order_column.asc(), Employee.id.asc())
    elif spec.dataset == "payroll":
        query, count_query = _payroll_query(company_ids, company_id, q, status_filter, year, period)
        payroll_sort = {"period": Payroll.period, "status": Payroll.status, "amount": Payroll.total_net, "net": Payroll.total_net, "gross": Payroll.total_gross}
        if sort in payroll_sort:
            order_column = payroll_sort[sort]
            query = query.order_by(order_column.asc() if sort in {"period", "status"} else order_column.desc(), Payroll.id.desc())
        else:
            query = query.order_by(Payroll.period.desc(), Payroll.id.desc())
    elif spec.dataset in {"payroll_line", "family_allowance"}:
        query, count_query = _payroll_line_query(company_ids, company_id, employee_id, q, status_filter, year, period, spec)
        line_sort = {"period": Payroll.period, "employee": Employee.full_name, "name": Employee.full_name, "amount": PayrollLine.net, "net": PayrollLine.net, "gross": PayrollLine.gross, "status": Payroll.status}
        if spec.dataset == "family_allowance":
            line_sort["amount"] = PayrollComplianceDetail.family_allowance
        if sort in line_sort:
            order_column = line_sort[sort]
            query = query.order_by(order_column.asc() if sort in {"period", "employee", "name", "status"} else order_column.desc(), PayrollLine.id.desc())
        else:
            query = query.order_by(Payroll.period.desc(), Employee.full_name.asc(), PayrollLine.id.desc())
    elif spec.dataset == "calculation":
        query, count_query = _calculation_query(company_ids, company_id, employee_id, q, status_filter, year, spec)
        calc_sort = {"date": CalculationRecord.created_at, "amount": CalculationRecord.amount, "type": CalculationRecord.calculation_type, "status": CalculationRecord.status, "period": CalculationRecord.reference_period}
        if sort in calc_sort:
            order_column = calc_sort[sort]
            query = query.order_by(order_column.asc() if sort in {"type", "status", "period"} else order_column.desc(), CalculationRecord.id.desc())
        else:
            query = query.order_by(CalculationRecord.created_at.desc(), CalculationRecord.id.desc())
    else:
        query, count_query = _related_query(spec.dataset, company_ids, company_id, employee_id, q, status_filter, year, period)
        if spec.dataset == "aguinaldo":
            sort_map = {"period": Aguinaldo.year, "year": Aguinaldo.year, "amount": Aguinaldo.calculated_amount, "employee": Employee.full_name, "name": Employee.full_name, "status": Aguinaldo.status, "date": Aguinaldo.created_at}
            if sort in sort_map:
                order_column = sort_map[sort]
                query = query.order_by(order_column.asc() if sort in {"period", "year", "employee", "name", "status"} else order_column.desc())
            else:
                query = query.order_by(Aguinaldo.year.desc(), Employee.full_name.asc())
        elif spec.dataset == "vacation":
            sort_map = {"period": Vacation.period_year, "year": Vacation.period_year, "employee": Employee.full_name, "name": Employee.full_name, "status": Vacation.status, "date": Vacation.start_date}
            if sort in sort_map:
                order_column = sort_map[sort]
                query = query.order_by(order_column.asc() if sort in {"employee", "name", "status"} else order_column.desc())
            else:
                query = query.order_by(Vacation.period_year.desc(), Employee.full_name.asc())
        elif spec.dataset == "medical_exam":
            sort_map = {"date": MedicalExam.exam_date, "employee": Employee.full_name, "name": Employee.full_name, "status": MedicalExam.result, "amount": MedicalExam.expires_on}
            if sort in sort_map:
                order_column = sort_map[sort]
                query = query.order_by(order_column.asc() if sort in {"employee", "name", "status"} else order_column.desc())
            else:
                query = query.order_by(MedicalExam.exam_date.desc(), MedicalExam.id.desc())
        elif spec.dataset == "operation":
            if spec.variant == "hours":
                query = query.where(or_(LaborOperation.operation_type.ilike("%hour%"), LaborOperation.operation_type.ilike("%hora%"), LaborOperation.operation_type.ilike("%overtime%"), LaborOperation.operation_type.ilike("%hora_extra%")))
            elif spec.variant == "advance":
                query = query.where(or_(LaborOperation.operation_type.ilike("%advance%"), LaborOperation.operation_type.ilike("%anticipo%")))
            elif spec.variant == "communication":
                query = query.where(or_(LaborOperation.operation_type.ilike("%telegram%"), LaborOperation.operation_type.ilike("%communication%"), LaborOperation.operation_type.ilike("%comunic%")))
            sort_map = {"date": LaborOperation.effective_date, "amount": LaborOperation.amount, "status": LaborOperation.status, "name": Employee.full_name}
            if sort in sort_map:
                order_column = sort_map[sort]
                query = query.order_by(order_column.asc() if sort in {"status", "name"} else order_column.desc(), LaborOperation.id.desc())
            else:
                query = query.order_by(LaborOperation.effective_date.desc(), LaborOperation.id.desc())
        elif spec.dataset == "certificate":
            sort_map = {"date": GeneratedCertificate.issue_date, "status": GeneratedCertificate.status, "name": GeneratedCertificate.employee_name_snapshot}
            if sort in sort_map:
                order_column = sort_map[sort]
                query = query.order_by(order_column.asc() if sort in {"status", "name"} else order_column.desc())
            else:
                query = query.order_by(GeneratedCertificate.issue_date.desc(), GeneratedCertificate.id.desc())
        elif spec.dataset == "employee_event":
            sort_map = {"date": EmployeeEvent.effective_date, "name": Employee.full_name, "status": EmployeeEvent.event_type}
            if sort in sort_map:
                order_column = sort_map[sort]
                query = query.order_by(order_column.asc() if sort in {"name", "status"} else order_column.desc())
            else:
                query = query.order_by(EmployeeEvent.effective_date.desc(), EmployeeEvent.id.desc())
        elif spec.dataset == "deadline":
            sort_map = {"date": LaborDeadline.due_date, "period": LaborDeadline.due_date, "status": LaborDeadline.status, "name": Company.legal_name}
            if sort in sort_map:
                order_column = sort_map[sort]
                query = query.order_by(order_column.asc(), LaborDeadline.id.desc())
            else:
                query = query.order_by(LaborDeadline.due_date.asc(), LaborDeadline.id.desc())
    # Count the fully filtered query itself so every joined filter is included in the total.
    total_query = select(func.count()).select_from(query.order_by(None).subquery())
    total = int(db.scalar(total_query) or 0)
    values = list(db.execute(query.offset((page - 1) * page_size).limit(page_size)).all())
    records = []
    for value in values:
        if spec.dataset in {"company", "employee", "payroll"}:
            record = _row(spec, value[0])
        else:
            record = _row(spec, value)
        records.append(record)
    return {
        "records": records,
        "total": total,
        "page": page,
        "page_size": page_size,
        "pages": (total + page_size - 1) // page_size,
        "columns": report_columns(spec),
        "filters": {"company_id": company_id, "employee_id": employee_id, "q": q, "status": status_filter, "year": year, "period": period, "sort": sort or ""},
    }


def _xlsx_safe_cell(value: Any) -> Any:
    """Store user-controlled strings as text, never as executable spreadsheet formulas."""
    if not isinstance(value, str):
        return value
    trimmed = value.lstrip(" \t\r\n")
    if value.startswith(("\t", "\r", "\n")) or (trimmed and trimmed[0] in {"=", "+", "-", "@"}):
        return "'" + value
    return value

def make_xlsx(spec: ReportSpec, rows: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Informe"
    sheet.append([spec.label])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(report_columns(spec))))
    sheet["A1"].font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="05256D")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.append(report_columns(spec))
    for cell in sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="163B78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 28
    for row in rows:
        sheet.append([_xlsx_safe_cell(row.get(column, "")) for column in report_columns(spec)])
    for row in sheet.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, int) and "(Gs.)" in str(sheet.cell(2, cell.column).value):
                cell.number_format = '#,##0;[Red]-#,##0'
    for idx, column in enumerate(report_columns(spec), start=1):
        values = [str(sheet.cell(row, idx).value or "") for row in range(2, min(sheet.max_row, 102) + 1)]
        width = max([len(column)] + [min(42, len(value)) for value in values])
        sheet.column_dimensions[get_column_letter(idx)].width = max(12, min(38, width + 2))
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(max(1, len(report_columns(spec))))}{sheet.max_row}"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_pdf(spec: ReportSpec, rows: list[dict[str, Any]], company_name: str = "") -> bytes:
    output = io.BytesIO()
    page = landscape(A4) if len(report_columns(spec)) > 7 else A4
    doc = SimpleDocTemplate(output, pagesize=page, rightMargin=10 * mm, leftMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm, title=spec.label)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ReportTitle", parent=styles["Title"], alignment=TA_LEFT, textColor=colors.HexColor("#05256D"), fontName="Helvetica-Bold", fontSize=15, leading=18, spaceAfter=4)
    small_style = ParagraphStyle("ReportSmall", parent=styles["Normal"], fontSize=7, leading=8, wordWrap="CJK")
    elements = [Paragraph(spec.label, title_style), Paragraph(xml_escape(f"Digit Laboral · Emitido {date.today().strftime('%d/%m/%Y')}" + (f" · {company_name}" if company_name else "")), small_style), Spacer(1, 4 * mm)]
    if spec.note:
        elements.extend([Paragraph(spec.note, small_style), Spacer(1, 3 * mm)])
    headers = [Paragraph(str(column), small_style) for column in report_columns(spec)]
    data = [headers]
    for row in rows:
        data.append([Paragraph(str(row.get(column, "") if row.get(column) is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>"), small_style) for column in report_columns(spec)])
    width = page[0] - 20 * mm
    col_widths = [width / max(1, len(report_columns(spec)))] * len(report_columns(spec))
    table = Table(data, colWidths=col_widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#05256D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CFD8E5")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5FA")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(table)
    doc.build(elements)
    return output.getvalue()
