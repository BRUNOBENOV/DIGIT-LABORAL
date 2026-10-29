from __future__ import annotations

import hashlib
import json
import re
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from .form_utils import clean_text
from .models import Branch, BranchComplianceProfile, CalculationRecord, Company, Employee


REOP_COMMUNICATIONS_URL = "https://www.mtess.gov.py/?page_id=31365"
REOP_FORMS_URL = "https://www.mtess.gov.py/?page_id=24021"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _json_dict(raw: str | None) -> dict[str, Any]:
    try:
        parsed = json.loads(raw or "{}")
        return parsed if isinstance(parsed, dict) else {}
    except (json.JSONDecodeError, TypeError):
        return {}


def _first_text(*values: Any, limit: int = 240) -> str:
    for value in values:
        if value is None:
            continue
        text = clean_text(str(value), max_length=limit)
        if text:
            return text
    return ""


def _branch_patronal(db: Session, company: Company, employee: Employee) -> tuple[str, str]:
    if not employee.branch_id:
        return company.mtess_employer_number or "", "Casa matriz"
    branch = db.get(Branch, employee.branch_id)
    profile = db.scalar(
        select(BranchComplianceProfile).where(BranchComplianceProfile.branch_id == employee.branch_id)
    )
    patronal = (
        (profile.mtess_employer_number if profile else "")
        or company.mtess_employer_number
        or ""
    )
    return patronal, branch.name if branch and branch.company_id == company.id else "Casa matriz"


def build_reop_salary_payload(
    db: Session,
    item: CalculationRecord,
    company: Company,
    employee: Employee,
    model: dict[str, Any],
    *,
    payment_date: str = "",
    payment_method: str = "",
    days_worked: str = "",
    hours_worked: str = "",
    pieces_or_tasks: str = "",
) -> dict[str, Any]:
    """Create an immutable handoff snapshot from a saved calculation; never recalculates wages."""
    inputs = _json_dict(item.input_json)
    result = _json_dict(item.result_json)
    identity = result.get("identity") if isinstance(result.get("identity"), dict) else {}
    employer_patronal, branch_name = _branch_patronal(db, company, employee)

    earnings = model.get("earnings", []) if isinstance(model.get("earnings"), list) else []
    deductions = model.get("deductions", []) if isinstance(model.get("deductions"), list) else []
    source_digest = hashlib.sha256(
        (item.input_json or "{}").encode("utf-8")
        + b"\n"
        + (item.result_json or "{}").encode("utf-8")
    ).hexdigest()

    return {
        "schema_version": 1,
        "integration_mode": "manual_assisted",
        "calculation_id": item.id,
        "calculation_status": item.status or "",
        "calculation_type": item.calculation_type or "",
        "period": model.get("period") or item.reference_period or "",
        "period_start": model["period_start"].isoformat() if model.get("period_start") else "",
        "period_end": model["period_end"].isoformat() if model.get("period_end") else "",
        "source_snapshot_sha256": source_digest,
        "company_id": company.id,
        "employer_name": _first_text(model.get("employer_name"), company.legal_name),
        "employer_ruc": _first_text(model.get("employer_ruc"), company.ruc, limit=60),
        "mtess_employer_number": _first_text(employer_patronal, model.get("mtess_employer_number"), limit=80),
        "branch_name": _first_text(branch_name, limit=160),
        "employer_address": _first_text(model.get("employer_address"), company.address),
        "employee_id": employee.id,
        "employee_name": _first_text(model.get("employee_name"), employee.full_name),
        "employee_ci": _first_text(model.get("employee_ci"), employee.document_number, limit=60),
        "employee_position": _first_text(model.get("employee_position"), employee.position, limit=160),
        "payment_method": _first_text(
            payment_method, model.get("payment_method"),
            identity.get("payment_method"), inputs.get("payment_method"), limit=100,
        ),
        "payment_date": _first_text(
            payment_date, result.get("payment_date"), result.get("paid_on"),
            inputs.get("payment_date"), inputs.get("paid_on"), limit=30,
        ),
        "days_worked": _first_text(
            days_worked, result.get("days_worked"), inputs.get("days_worked"),
            inputs.get("worked_days"), limit=30,
        ),
        "hours_worked": _first_text(
            hours_worked, result.get("hours_worked"), inputs.get("hours_worked"),
            inputs.get("worked_hours"), limit=30,
        ),
        "pieces_or_tasks": _first_text(
            pieces_or_tasks, result.get("pieces_or_tasks"), inputs.get("pieces_or_tasks"),
            inputs.get("tasks_count"), limit=120,
        ),
        "earnings": [
            {
                "label": _first_text(row.get("label"), "Concepto de ingreso", limit=160),
                "amount": int(row.get("amount") or 0),
                "formula": _first_text(row.get("formula"), limit=240),
                "legal": _first_text(row.get("legal"), limit=160),
            }
            for row in earnings if isinstance(row, dict)
        ],
        "deductions": [
            {
                "label": _first_text(row.get("label"), "Concepto de descuento", limit=160),
                "amount": int(row.get("amount") or 0),
                "formula": _first_text(row.get("formula"), limit=240),
                "legal": _first_text(row.get("legal"), limit=160),
            }
            for row in deductions if isinstance(row, dict)
        ],
        "gross": int(model.get("gross") or 0),
        "discounts": int(model.get("discounts") or 0),
        "net": int(model.get("net") or 0),
        "calculation_warnings": [
            clean_text(str(value), max_length=350)
            for value in (model.get("warnings") or [])
            if value and str(value).strip()
        ],
        "manual_review_note": (
            "Borrador de apoyo para la carga humana en REOP. No es una plantilla oficial de importación "
            "ni representa un envío o recibo oficial del MTESS."
        ),
    }


def validate_reop_salary_payload(payload: dict[str, Any]) -> list[str]:
    """Validate the data currently available; this is not a legal certification by MTESS."""
    errors: list[str] = []
    required_labels = {
        "employer_name": "Razón social del empleador",
        "employer_ruc": "RUC del empleador",
        "mtess_employer_number": "Número patronal MTESS",
        "employee_name": "Nombre del funcionario",
        "employee_ci": "Cédula del funcionario",
        "period_start": "Inicio del período de pago",
        "period_end": "Fin del período de pago",
        "payment_method": "Forma de pago",
        "payment_date": "Fecha de pago",
        "days_worked": "Cantidad de días trabajados",
        "hours_worked": "Cantidad de horas trabajadas",
        "pieces_or_tasks": "Cantidad de piezas/tareas o indicación «No aplica»",
    }
    for key, label in required_labels.items():
        if not str(payload.get(key) or "").strip():
            errors.append(f"Falta completar: {label}.")

    payment_date = str(payload.get("payment_date") or "").strip()
    if payment_date:
        try:
            from datetime import date
            date.fromisoformat(payment_date)
        except ValueError:
            errors.append("La fecha de pago debe tener formato AAAA-MM-DD.")

    for key, label, maximum in (
        ("days_worked", "Cantidad de días trabajados", Decimal("31")),
        ("hours_worked", "Cantidad de horas trabajadas", Decimal("744")),
    ):
        raw = str(payload.get(key) or "").strip()
        if raw:
            try:
                numeric = Decimal(raw)
                if not numeric.is_finite() or numeric < 0 or numeric > maximum:
                    errors.append(f"{label}: el valor debe estar entre 0 y {maximum}.")
            except InvalidOperation:
                errors.append(f"{label}: ingresá un número válido.")

    try:
        from datetime import date
        start = date.fromisoformat(str(payload.get("period_start") or ""))
        end = date.fromisoformat(str(payload.get("period_end") or ""))
        if end < start:
            errors.append("El fin del período no puede ser anterior al inicio.")
    except ValueError:
        if payload.get("period_start") or payload.get("period_end"):
            errors.append("El período debe contener fechas válidas en formato AAAA-MM-DD.")

    if not payload.get("earnings"):
        errors.append("El cálculo guardado no contiene conceptos de ingresos para comunicar.")
    if int(payload.get("gross") or 0) != sum(int(row.get("amount") or 0) for row in payload.get("earnings", [])):
        errors.append("La suma de conceptos de ingresos no coincide con el bruto guardado.")
    if int(payload.get("discounts") or 0) != sum(int(row.get("amount") or 0) for row in payload.get("deductions", [])):
        errors.append("La suma de conceptos de descuentos no coincide con el total guardado.")
    if int(payload.get("gross") or 0) - int(payload.get("discounts") or 0) != int(payload.get("net") or 0):
        errors.append("El bruto menos los descuentos no coincide con el neto guardado.")

    blocked_fragments = (
        "El período no tiene formato AAAA-MM válido",
        "Falta el nombre patronal o el RUC",
        "Falta la identidad documental del funcionario",
        "La suma de conceptos de haberes no coincide",
        "La suma de conceptos de descuentos no coincide",
        "El bruto menos los descuentos no coincide",
    )
    for warning in payload.get("calculation_warnings", []):
        if any(fragment in str(warning) for fragment in blocked_fragments):
            errors.append(f"Validación del cálculo de origen: {warning}")

    return list(dict.fromkeys(errors))


def _safe_excel_text(value: Any) -> str:
    text = "" if value is None else str(value)
    # Prevent untrusted labels/names from being interpreted as spreadsheet formulas.
    if text.lstrip(" \t\r\n")[:1] in {"=", "+", "-", "@"}:
        return "'" + text
    return text


def _style_sheet(ws) -> None:  # noqa: ANN001
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="173B86")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for column in ws.columns:
        length = max((len(str(cell.value or "")) for cell in column), default=12)
        ws.column_dimensions[column[0].column_letter].width = min(48, max(14, length + 2))


def build_reop_salary_handoff_workbook(payload: dict[str, Any]) -> bytes:
    """Build a review workbook; deliberately not shaped as an official REOP import template."""
    wb = Workbook()
    summary = wb.active
    summary.title = "DATOS_PARA_CARGA"
    summary.append(["Campo", "Valor"])
    summary_rows = [
        ("Tipo de documento", "Borrador de apoyo para carga manual en REOP"),
        ("ID del cálculo DIGIT-LABORAL", payload.get("calculation_id")),
        ("Hash del cálculo de origen (SHA-256)", payload.get("source_snapshot_sha256")),
        ("Razón social", payload.get("employer_name")),
        ("RUC", payload.get("employer_ruc")),
        ("Número patronal MTESS", payload.get("mtess_employer_number")),
        ("Establecimiento", payload.get("branch_name")),
        ("Nombre del trabajador", payload.get("employee_name")),
        ("Cédula", payload.get("employee_ci")),
        ("Cargo", payload.get("employee_position")),
        ("Período desde", payload.get("period_start")),
        ("Período hasta", payload.get("period_end")),
        ("Forma de pago", payload.get("payment_method")),
        ("Fecha de pago", payload.get("payment_date")),
        ("Días trabajados", payload.get("days_worked")),
        ("Horas trabajadas", payload.get("hours_worked")),
        ("Piezas/tareas", payload.get("pieces_or_tasks")),
        ("Total ingresos guardado (Gs.)", payload.get("gross")),
        ("Total descuentos guardado (Gs.)", payload.get("discounts")),
        ("Neto guardado (Gs.)", payload.get("net")),
        ("Estado del cálculo de origen", payload.get("calculation_status")),
        ("Advertencia", "NO es una planilla oficial de importación. No cargar este archivo en Importar Planilla sin validación expresa del MTESS."),
    ]
    for label, value in summary_rows:
        summary.append([_safe_excel_text(label), _safe_excel_text(value) if isinstance(value, (str, type(None))) else value])
    _style_sheet(summary)

    for sheet_name, key in (("INGRESOS", "earnings"), ("DESCUENTOS", "deductions")):
        ws = wb.create_sheet(sheet_name)
        ws.append(["Concepto guardado", "Importe (Gs.)", "Fórmula / detalle", "Referencia legal registrada"])
        for row in payload.get(key, []):
            ws.append([
                _safe_excel_text(row.get("label")),
                int(row.get("amount") or 0),
                _safe_excel_text(row.get("formula")),
                _safe_excel_text(row.get("legal")),
            ])
        _style_sheet(ws)

    validations = wb.create_sheet("VALIDACIONES")
    validations.append(["Resultado", "Detalle"])
    errors = validate_reop_salary_payload(payload)
    if errors:
        for error in errors:
            validations.append(["PENDIENTE", _safe_excel_text(error)])
    else:
        validations.append(["LISTO PARA REVISIÓN HUMANA", "No se detectaron faltantes en las comprobaciones técnicas locales."])
    validations.append([
        "LÍMITE DEL SISTEMA",
        "La validación local no equivale a aceptación del MTESS. La presentación y la generación del recibo oficial se realizan en REOP.",
    ])
    _style_sheet(validations)

    trace = wb.create_sheet("TRAZABILIDAD")
    trace.append(["Campo", "Valor"])
    for label, value in (
        ("Plataforma", "DIGIT-LABORAL"),
        ("Autoridad de destino", "MTESS / REOP"),
        ("Tipo de comunicación", "Liquidación salarial"),
        ("Modalidad", "Preparación y presentación manual asistida"),
        ("Cálculo de origen", payload.get("calculation_id")),
        ("Hash SHA-256 del origen", payload.get("source_snapshot_sha256")),
        ("Fuente oficial de instrucciones", REOP_COMMUNICATIONS_URL),
        ("Especificación pública de API de terceros", "No localizada en la documentación oficial consultada; transmisión directa deshabilitada."),
        ("Importante", "Este archivo no es un recibo oficial, no acredita envío y no debe confundirse con la planilla oficial vigente."),
    ):
        trace.append([_safe_excel_text(label), _safe_excel_text(value) if isinstance(value, (str, type(None))) else value])
    _style_sheet(trace)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
